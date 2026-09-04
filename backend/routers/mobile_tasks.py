"""组长、组员使用的手机任务首页和卡片式处理接口。"""

from __future__ import annotations

from datetime import datetime
import json
import logging
from io import BytesIO
import time
from typing import Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from database import db_manager, get_db
from deps import get_current_user, require_permission
from routers.query import (
    _enabled_spreadsheets,
    _managed_column_metadata,
    _source_ready,
    _writeback_enabled,
    queue_source_fields,
)
from routers.workflow import _can_view_ticket as workflow_ticket_access
from services.business_time import get_business_date
from services.data_scope import community_names_for_scopes
from services.online_edit_permissions import (
    effective_view_communities,
    inspector_option_context,
    row_edit_capabilities,
)
from services.online_source import (
    active_source_sql_filter,
    json_value,
    rebuild_projection,
    rebuild_projection_rows,
)
from services.address_matching import MATCHER_VERSION
from services.address_match_feedback import record_feedback_confirmation
from services.online_local_writeback import (
    apply_local_system_changes,
    launch_local_change_processing,
    load_local_changes,
    local_sync_state,
    overlay_local_values,
    resolve_source_conflict,
)
from services.parsers import get_parser
from services.permissions import (
    ONLINE_RAW_EDIT,
    ONLINE_RAW_VIEW,
    ONLINE_TASK_MANAGE,
    QMF_REGISTRATION_EXECUTE,
    WORKFLOW_ATTACHMENT_VIEW,
    has_permission,
)
from services.report_overview import SUMMARY_TYPE, get_online_overview
from services.qmf_registration import (
    MODEL_THREE_PARSER,
    preview_capability,
    registration_capability,
)
from services.qmf_status import normalize_qmf_status_result
from services.qmf_config import load_qmf_config
from services.qmf_runs import WRITE_STEP_KEYS, parse_steps, utc_text
from services.task_workflow import (
    MOBILE_TASK_TYPES,
    PHONE_FIELD_ALIASES,
    SUMMARY_TASK_TYPES,
    TASK_WORKFLOWS,
)
from services.task_graph import online_task_blocked
from services.task_assignment_responsibility import (
    capture_first_assignment,
    record_internal_transfer,
)
from services.unverifiable_review import (
    ACTIVE_STATES as UNVERIFIABLE_ACTIVE_STATES,
    DEEP_EXTENSION,
    DEEP_PENDING,
    FINAL_UNVERIFIABLE,
    INITIAL_EXTENSION,
    INITIAL_PENDING,
    STATE_LABELS as UNVERIFIABLE_STATE_LABELS,
    UNVERIFIABLE_REVIEW_TYPES,
    apply_decision,
    prepare_decision,
    review_events_for_flow,
    review_flows_by_rows,
    supports_unverifiable_review,
)
from services.audit import record_admin_audit, request_audit_fields
from services.xlsx_export import XLSX_MEDIA_TYPE, build_xlsx
from config import settings
from services.local_source import local_data_source_enabled, local_row_hash
from services.watch_matching import task_watch_payload
from services.residence_platform import ResidencePlatformError
from services.residence_status_scan import (
    residence_detail_for_values,
    residence_status_by_rows,
    wake_residence_lookup_scheduler,
)
from services.task_registration import (
    REGISTRATION_TASK_TYPES,
    cancel_registration_link,
    is_registration_task,
    migrate_registration_link,
    refresh_registration_source_context_after_writeback,
    registration_links_by_rows,
    select_registration_property,
    validate_registration_property,
    record_registration_event,
)
from services.residence_platform_config import load_residence_config
from services.residence_status_scan import _load_current_target, _lookup_registration_address_target


router = APIRouter(prefix="/api/mobile-tasks", tags=["手机任务工作台"])
logger = logging.getLogger(__name__)
FlowScope = Literal["mine", "community", "all"]
TaskStatus = Literal[
    "pending",
    "unchecked",
    "checked",
    "review",
    "registration_review",
    "completed",
    "all",
]
ReviewStage = Literal[
    "all", "waiting_analysis", "analyzed",
    "initial_pending", "initial_extension", "deep_pending", "deep_extension",
]
Priority = Literal[
    "all",
    "analyzed",
    "source_exception",
    "pending_sync",
    "ordinary",
    "waiting_analysis",
    "completed",
]
SortMode = Literal[
    "priority",
    "address_asc",
    "identity_asc",
    "updated_desc",
    "updated_asc",
]
AssignmentMode = Literal["single", "balanced"]
QmfFeedbackState = Literal[
    "not_scanned",
    "stale",
    "pending",
    "completed_match",
    "completed_mismatch",
    "not_found",
    "non_jurisdiction",
    "error",
]
EMPTY_FILTER_VALUE = "__empty__"
MAX_BULK_ASSIGNMENT_TASKS = 2000
MAX_BULK_ASSIGNMENT_CHUNK = 100


async def _task_source_ready(cur, parser_type: str) -> bool:
    """Local mode is ready when the local projection exists, not when a Tencent
    spreadsheet configuration exists.  An empty local table is still a valid
    ready state and should render an empty task pool rather than a sync error.
    """
    if local_data_source_enabled():
        await cur.execute(
            "SELECT 1 FROM _online_source_rows AS source "
            "WHERE source.parser_type=%s AND source.archived_at IS NULL "
            f"{active_source_sql_filter(parser_type, 'source')} LIMIT 1",
            (parser_type,),
        )
        await cur.fetchone()
        return True
    return await _source_ready(cur, await _enabled_spreadsheets(cur, parser_type))


async def _active_source_counts(cur, parser_type: str, row_keys: list[str]) -> dict[str, int]:
    if not local_data_source_enabled() or not row_keys:
        return {}
    placeholders = ",".join(["%s"] * len(row_keys))
    await cur.execute(
        "SELECT row_key, COUNT(*) FROM _online_source_rows AS source "
        "WHERE source.parser_type=%s AND source.row_key IN (" + placeholders + ") "
        "AND source.archived_at IS NULL "
        f"{active_source_sql_filter(parser_type, 'source')} GROUP BY row_key",
        [parser_type, *row_keys],
    )
    return {
        str(row_key): int(count or 0)
        for row_key, count in await cur.fetchall()
    }


async def _qmf_registration_state(
    conn,
    *,
    parser_type: str,
    sources: list[dict],
    user: dict,
) -> tuple[dict | None, dict | None]:
    """Return a private resumable run plus the public successful summary.

    Only users with the全民防 registration permission receive prepared,
    executing, failed or uncertain details. Other task viewers can only see
    that the external feedback succeeded and when it completed.
    """
    if parser_type != MODEL_THREE_PARSER or not sources:
        return None, None
    source_ids = [int(item["id"]) for item in sources]
    placeholders = ",".join(["%s"] * len(source_ids))
    latest_run = None
    async with conn.cursor() as cur:
        if has_permission(user, QMF_REGISTRATION_EXECUTE):
            await cur.execute(
                "UPDATE _qmf_registration_runs "
                "SET status='expired', result_code='prepare_expired' "
                f"WHERE parser_type=%s AND source_id IN ({placeholders}) "
                "AND requested_by=%s AND status='prepared' "
                "AND expires_at<=UTC_TIMESTAMP()",
                (parser_type, *source_ids, int(user["id"])),
            )
            await cur.execute(
                "SELECT id, source_id, expected_revision, status, steps_json, "
                "result_code, tencent_marker_status, tencent_marker_error, "
                "prepared_at, expires_at, execution_started_at, completed_at, "
                "created_at, updated_at "
                f"FROM _qmf_registration_runs WHERE parser_type=%s "
                f"AND source_id IN ({placeholders}) AND requested_by=%s "
                "ORDER BY id DESC LIMIT 1",
                (parser_type, *source_ids, int(user["id"])),
            )
            row = await cur.fetchone()
            if row:
                status = str(row[3] or "")
                marker_status = str(row[6] or "not_started")
                steps = parse_steps(row[4])
                write_progress = any(
                    item["key"] in WRITE_STEP_KEYS
                    and item["status"] in {"sending", "succeeded", "uncertain"}
                    for item in steps
                )
                latest_run = {
                    "id": int(row[0]),
                    "parser_type": parser_type,
                    "source_id": int(row[1]),
                    "expected_revision": int(row[2]),
                    "status": status,
                    "steps": steps,
                    "result_code": str(row[5] or ""),
                    "photo": {"sha256": "", "mime_type": "", "size_bytes": 0},
                    "tencent_marker_status": marker_status,
                    "tencent_marker_error": str(row[7] or ""),
                    "prepared_at": utc_text(row[8]),
                    "expires_at": utc_text(row[9]),
                    "execution_started_at": utc_text(row[10]),
                    "completed_at": utc_text(row[11]),
                    "created_at": utc_text(row[12]),
                    "updated_at": utc_text(row[13]),
                    "can_execute": status == "prepared",
                    "can_reprepare": status in {"failed", "uncertain"}
                    and not write_progress,
                    "can_retry_marker": status == "succeeded" and marker_status in {
                        "not_started", "pending", "conflict", "failed",
                    },
                }
        await cur.execute(
            "SELECT id, completed_at, tencent_marker_status "
            f"FROM _qmf_registration_runs WHERE parser_type=%s "
            f"AND source_id IN ({placeholders}) AND status='succeeded' "
            "ORDER BY id DESC LIMIT 1",
            (parser_type, *source_ids),
        )
        succeeded = await cur.fetchone()
    feedback = None
    if succeeded:
        feedback = {
            "run_id": int(succeeded[0]),
            "status": "succeeded",
            "completed_at": utc_text(succeeded[1]),
            "tencent_marker_status": str(succeeded[2] or "not_started"),
        }
    return latest_run, feedback


async def _task_photo_results(user: dict, parser_type: str, row_key: str) -> list[dict]:
    """返回与当前任务精确关联的已完成照片工单附件摘要。"""
    if (
        not settings.WORKFLOW_FEATURE_ENABLED
        or not has_permission(user, WORKFLOW_ATTACHMENT_VIEW)
    ):
        return []
    try:
        pool = db_manager.get_pool("workflow")
    except ValueError:
        return []
    results: list[dict] = []
    async with pool.acquire() as workflow_conn:
        async with workflow_conn.cursor() as cur:
            await cur.execute(
                "SELECT order_row.id, order_row.ticket_no "
                "FROM photo_request_details detail "
                "JOIN work_orders order_row ON order_row.id=detail.work_order_id "
                "WHERE detail.external_origin='platform_task' "
                "AND detail.source_parser_type=%s AND detail.source_row_key=%s "
                "AND detail.result_status='found' "
                "AND order_row.type_code='photo_request' "
                "AND order_row.status IN ('approved','completed') "
                "ORDER BY order_row.completed_at DESC, order_row.id DESC",
                (parser_type, row_key),
            )
            tickets = await cur.fetchall()
            for ticket_id, ticket_no in tickets:
                normalized_ticket_id = int(ticket_id)
                try:
                    await workflow_ticket_access(cur, normalized_ticket_id, user)
                except HTTPException as exc:
                    if exc.status_code in {403, 404}:
                        continue
                    raise
                await cur.execute(
                    "SELECT file_id, original_name, mime_type, size_bytes "
                    "FROM work_order_attachments "
                    "WHERE work_order_id=%s AND deleted_at IS NULL "
                    "AND mime_type LIKE 'image/%%' ORDER BY id DESC",
                    (normalized_ticket_id,),
                )
                attachments = [
                    {
                        "file_id": str(file_id),
                        "original_name": str(original_name or "照片"),
                        "mime_type": str(mime_type or "application/octet-stream"),
                        "size_bytes": int(size_bytes or 0),
                    }
                    for file_id, original_name, mime_type, size_bytes in await cur.fetchall()
                ]
                if attachments:
                    results.append({
                        "ticket_id": normalized_ticket_id,
                        "ticket_no": str(ticket_no or ""),
                        "attachments": attachments,
                    })
    return results


async def _task_photo_fetched_rows(
    user: dict,
    parser_type: str,
    row_keys: list[str],
) -> set[str]:
    """Return task row keys with an accessible, completed photo result.

    The list view asks for several rows at once so photo status does not cause
    one workflow-database query per task card.  Ticket visibility still goes
    through the same access check used by the detail endpoint.
    """
    if (
        not row_keys
        or not settings.WORKFLOW_FEATURE_ENABLED
        or not has_permission(user, WORKFLOW_ATTACHMENT_VIEW)
    ):
        return set()
    try:
        pool = db_manager.get_pool("workflow")
    except ValueError:
        return set()
    placeholders = ",".join(["%s"] * len(row_keys))
    fetched: set[str] = set()
    async with pool.acquire() as workflow_conn:
        async with workflow_conn.cursor() as cur:
            await cur.execute(
                "SELECT DISTINCT detail.source_row_key, order_row.id "
                "FROM photo_request_details detail "
                "JOIN work_orders order_row ON order_row.id=detail.work_order_id "
                "JOIN work_order_attachments attachment "
                "ON attachment.work_order_id=order_row.id "
                "AND attachment.deleted_at IS NULL "
                "AND attachment.mime_type LIKE 'image/%%' "
                "WHERE detail.external_origin='platform_task' "
                "AND detail.source_parser_type=%s "
                f"AND detail.source_row_key IN ({placeholders}) "
                "AND detail.result_status='found' "
                "AND order_row.type_code='photo_request' "
                "AND order_row.status IN ('approved','completed')",
                [parser_type, *row_keys],
            )
            candidates = await cur.fetchall()
            for source_row_key, ticket_id in candidates:
                normalized_ticket_id = int(ticket_id)
                try:
                    await workflow_ticket_access(cur, normalized_ticket_id, user)
                except HTTPException as exc:
                    if exc.status_code in {403, 404}:
                        continue
                    raise
                fetched.add(str(source_row_key))
    return fetched
PRIORITY_KEYS = (
    "analyzed",
    "source_exception",
    "pending_sync",
    "ordinary",
    "waiting_analysis",
    "completed",
)
PRIORITY_LABELS = {
    "analyzed": "已研判",
    "source_exception": "来源异常",
    "pending_sync": "待同步",
    "ordinary": "普通待处理",
    "waiting_analysis": "等待研判",
    "completed": "已完成",
}
FLOW_TASK_ELEVATED_POSITIONS = {
    "片长",
    "基础管控",
    "中队长",
    "社区民警",
    "所队领导",
}


class TaskBatchUpdate(BaseModel):
    changes: dict[str, str]
    base_values: dict[str, str] = Field(default_factory=dict)
    expected_revision: int = Field(gt=0)
    registration_property_id: int | None = Field(default=None, gt=0)
    registration_property_version: int | None = Field(default=None, gt=0)


class UnverifiableDecision(BaseModel):
    """Structured two-stage decision for an unresolved task."""

    stage: Literal["initial_pending", "deep_pending"]
    outcome: Literal["success", "failure"]
    opinion: str = Field(min_length=1, max_length=2000)
    flow_version: int = Field(gt=0)
    expected_revision: int = Field(gt=0)
    expected_row_hash: str = Field(min_length=1, max_length=128)


class RegistrationManualConfirm(BaseModel):
    reason: Literal["address_mismatch", "address_ambiguous"]
    note: str = Field(default="", max_length=500)
    expected_revision: int = Field(gt=0)


class AddressMatchConfirm(BaseModel):
    small_community_id: int = Field(gt=0)


class AddressMatchConflictResolution(BaseModel):
    """管理员确认冲突候选后，将任务社区改为候选正式社区并重新匹配。"""

    small_community_id: int = Field(gt=0)
    expected_revision: int = Field(gt=0)
    expected_row_hash: str = Field(min_length=1, max_length=128)


class SyncConflictResolution(BaseModel):
    choice: Literal["platform", "tencent"]
    fields: list[str] = Field(min_length=1, max_length=5)


class BulkAssignmentRequest(BaseModel):
    row_keys: list[str] = Field(min_length=1, max_length=MAX_BULK_ASSIGNMENT_CHUNK)
    inspector: str = Field(default="", max_length=100)
    mode: AssignmentMode = "single"
    balanced_offset: int = Field(default=0, ge=0)
    balanced_total: int = Field(default=0, ge=0, le=MAX_BULK_ASSIGNMENT_TASKS)


class CancelAssignmentsRequest(BaseModel):
    """按业务和可选社区撤销尚未完成任务的核查人分配。"""

    community: str = Field(default="", max_length=200)


class InternalTransferRequest(BaseModel):
    target_community: str = Field(min_length=1, max_length=200)
    target_leader: str = Field(default="", max_length=100)
    expected_row_key: str = Field(min_length=1, max_length=500)
    expected_revision: int = Field(gt=0)
    expected_row_hash: str = Field(min_length=1, max_length=128)


class TaskSearch(BaseModel):
    scope: FlowScope = "mine"
    status: TaskStatus = "pending"
    review_stage: ReviewStage = "all"
    communities: list[str] = Field(default_factory=list, max_length=50)
    small_communities: list[str] = Field(default_factory=list, max_length=50)
    match_status: list[str] = Field(default_factory=list, max_length=20)
    inspectors: list[str] = Field(default_factory=list, max_length=50)
    watch_categories: list[int] = Field(default_factory=list, max_length=50)
    qmf_feedback_states: list[QmfFeedbackState] = Field(
        default_factory=list, max_length=10
    )
    priority: Priority = "all"
    sort: SortMode = "priority"
    keyword: str = Field(default="", max_length=100)
    page: int = Field(default=1, ge=1)
    page_size: int = Field(default=20, ge=1, le=50)


class AnalysisTaskSearch(BaseModel):
    parser_types: list[str] = Field(min_length=1, max_length=20)
    scope: FlowScope = "all"
    review_stage: ReviewStage = "all"
    communities: list[str] = Field(default_factory=list, max_length=50)
    small_communities: list[str] = Field(default_factory=list, max_length=50)
    match_status: list[str] = Field(default_factory=list, max_length=20)
    inspectors: list[str] = Field(default_factory=list, max_length=50)
    watch_categories: list[int] = Field(default_factory=list, max_length=50)
    sort: SortMode = "priority"
    keyword: str = Field(default="", max_length=100)
    page: int = Field(default=1, ge=1)
    page_size: int = Field(default=20, ge=1, le=50)


class InlineEditorRequest(BaseModel):
    row_keys: list[str] = Field(min_length=1, max_length=50)


def _iso_utc(value) -> str | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.isoformat() + "Z"
    return str(value)


def _balanced_assignment_plan(
    row_keys: list[str],
    inspectors: list[str],
    *,
    total_count: int | None = None,
    start_index: int = 0,
) -> tuple[dict[str, str], dict[str, int]]:
    """把已筛出的任务按连续区段尽量平均分给在岗组员。

    连续区段比轮流交错更适合任务列表的地址排序：相邻地址尽量仍由同一
    名组员处理，同时每人的任务数最多相差一条。真正的候选人和任务资格
    仍由批量接口在数据库事务前后重新校验。
    """
    keys = list(dict.fromkeys(str(key).strip() for key in row_keys if str(key).strip()))
    members = list(dict.fromkeys(str(name).strip() for name in inspectors if str(name).strip()))
    if not keys or not members:
        return {}, {name: 0 for name in members}
    total = total_count if total_count is not None else len(keys)
    if total < len(keys) or start_index < 0 or start_index + len(keys) > total:
        return {}, {name: 0 for name in members}
    base, remainder = divmod(total, len(members))
    member_ranges: list[tuple[int, int, str]] = []
    range_start = 0
    for index, member in enumerate(members):
        size = base + (1 if index < remainder else 0)
        member_ranges.append((range_start, range_start + size, member))
        range_start += size
    plan: dict[str, str] = {}
    counts = {name: 0 for name in members}
    for local_index, key in enumerate(keys):
        global_index = start_index + local_index
        member = next(
            name
            for lower, upper, name in member_ranges
            if lower <= global_index < upper
        )
        plan[key] = member
        counts[member] += 1
    return plan, counts


def _bulk_assignment_result(
    *,
    updated: int,
    skipped: list[dict[str, str]],
    failures: list[dict[str, str]],
    inspector: str,
    mode: AssignmentMode,
    assignment_counts: dict[str, int],
    duration_ms: int = 0,
) -> dict:
    """Build mutually exclusive assignment outcome counts and details."""
    return {
        "updated": updated,
        "skipped": len(skipped),
        "failed": len(failures),
        "details": skipped,
        "failed_details": failures,
        "inspector": inspector if mode == "single" else "",
        "mode": mode,
        "assignment_counts": assignment_counts,
        "duration_ms": duration_ms,
    }


def require_flow_user(user: dict) -> tuple[str, str]:
    member = user.get("member") or {}
    position = str(member.get("position") or "").strip()
    name = str(member.get("name") or "").strip()
    communities = [
        str(value).strip()
        for value in user.get("community_names") or []
        if str(value).strip()
    ]
    if position not in {"组员", "组长"}:
        raise HTTPException(403, "手机任务工作台目前仅向组员和组长开放")
    if not name:
        raise HTTPException(403, "当前账号尚未关联有效人员")
    if len(communities) != 1:
        raise HTTPException(403, "当前人员必须配置一个有效社区部门")
    return name, communities[0]


def is_flow_task_admin(user: dict) -> bool:
    group_codes = {
        str(group.get("code") or "")
        for group in user.get("permission_groups") or []
        if isinstance(group, dict)
    }
    legacy_group = str((user.get("permission_group") or {}).get("code") or "")
    return (
        str(user.get("role") or "") in {"admin", "super_admin"}
        or bool(group_codes & {"admin", "super_admin"})
        or legacy_group in {"admin", "super_admin"}
    )


def is_flow_task_elevated(user: dict) -> bool:
    position = str((user.get("member") or {}).get("position") or "").strip()
    return (
        is_flow_task_admin(user)
        or has_permission(user, ONLINE_TASK_MANAGE)
        or position in FLOW_TASK_ELEVATED_POSITIONS
    )


REGISTRATION_CONFIRM_POSITIONS = {"基础管控", "中队长", "所队领导"}


def can_confirm_registration(user: dict) -> bool:
    position = str((user.get("member") or {}).get("position") or "").strip()
    return is_flow_task_admin(user) or position in REGISTRATION_CONFIRM_POSITIONS


def _registration_update_hooks(
    parser_type: str,
    data: TaskBatchUpdate,
    user: dict,
):
    """Build transaction hooks for an atomic result/property save."""
    if not is_registration_task(parser_type):
        if data.registration_property_id or data.registration_property_version:
            raise HTTPException(400, "该业务不支持拟登记房屋关联")
        return None, None, False
    if bool(data.registration_property_id) != bool(data.registration_property_version):
        raise HTTPException(422, "房屋编号和房屋版本必须同时提交")

    workflow = TASK_WORKFLOWS[parser_type]
    prepared: dict = {}

    async def transaction_prepare(*, cur, source, current_values, changes):
        current_result = str(current_values.get(workflow.result_field) or "").strip()
        after = dict(current_values)
        after.update(changes)
        result = str(after.get(workflow.result_field) or "").strip()
        submitted_result = (
            str(changes.get(workflow.result_field) or "").strip()
            if workflow.result_field in changes
            else None
        )
        result_changed = submitted_result is not None and submitted_result != current_result
        if submitted_result == "已登记" and current_result != "已登记":
            raise HTTPException(
                403,
                "不能直接选择已登记；请等待居住证自动比对或由有权人员复核确认",
            )
        links = await registration_links_by_rows(
            cur,
            parser_type,
            [str(source["row_key"])],
        )
        existing = links.get(str(source["row_key"]))
        prepared["existing"] = existing
        if result != "待登记":
            if data.registration_property_id:
                raise HTTPException(400, "只有待登记结果可以关联拟登记房屋")
            prepared["action"] = (
                "cancel"
                if current_result == "待登记" and result_changed
                else "preserve"
            )
            return {}

        if (
            existing
            and existing.get("status") == "legacy_completed"
            and current_result == "待登记"
            and not result_changed
            and not data.registration_property_id
        ):
            # Historical rows keep their completed compatibility state when
            # another field is edited.  They are not forced into the new
            # property-selection contract retroactively.
            prepared["action"] = "preserve"
            return {}

        prepared["action"] = "pending"

        if data.registration_property_id:
            try:
                property_row = await validate_registration_property(
                    cur,
                    property_id=int(data.registration_property_id),
                    expected_version=int(data.registration_property_version or 0),
                    task_community=get_parser(parser_type).community_value(after),
                )
            except ValueError as exc:
                raise HTTPException(409, str(exc)) from exc
            prepared["property"] = property_row
            return {"现住址": property_row["address"]}

        existing_usable = bool(
            existing
            and existing.get("property_id")
            and int(existing.get("source_id") or 0) == int(source["id"])
            and existing.get("status") not in {"cancelled", "legacy_completed"}
        )
        if not existing_usable:
            raise HTTPException(422, "选择待登记时必须同时选择唯一的拟登记房屋")
        return {}

    async def transaction_callback(
        *, cur, source, before, after, row_key_before, row_key_after, revision,
    ):
        del before
        action = prepared.get("action")
        if action == "preserve":
            return
        if action == "pending":
            await cur.execute(
                "SELECT COALESCE(identity_hmac,''),community "
                "FROM _online_source_projection WHERE parser_type=%s AND row_key=%s",
                (parser_type, row_key_after),
            )
            projection_context = await cur.fetchone()
            if not projection_context:
                raise HTTPException(409, "任务投影已变化，请刷新后重试")
            if row_key_before != row_key_after and not prepared.get("property"):
                try:
                    await migrate_registration_link(
                        cur,
                        parser_type=parser_type,
                        row_key_before=row_key_before,
                        row_key_after=row_key_after,
                        source_id=int(source["id"]),
                    )
                except ValueError as exc:
                    raise HTTPException(409, str(exc)) from exc
            property_row = prepared.get("property")
            if property_row:
                if row_key_before != row_key_after:
                    existing = prepared.get("existing")
                    if existing:
                        try:
                            await migrate_registration_link(
                                cur,
                                parser_type=parser_type,
                                row_key_before=row_key_before,
                                row_key_after=row_key_after,
                                source_id=int(source["id"]),
                            )
                        except ValueError as exc:
                            raise HTTPException(409, str(exc)) from exc
                await select_registration_property(
                    cur,
                    parser_type=parser_type,
                    row_key=row_key_after,
                    source_id=int(source["id"]),
                    property_id=int(property_row["id"]),
                    property_version=int(property_row["version"]),
                    source_revision=int(revision),
                    source_row_hash=local_row_hash(after),
                    identity_hmac=str(projection_context[0] or ""),
                    task_community=str(projection_context[1] or ""),
                    user_id=int(user.get("id")) if user.get("id") else None,
                )
            else:
                # A normal edit on an already-linked pending task advances the
                # local source revision and content hash in the same transaction.
                # Carry the link across that exact change so the scanner cannot
                # mistake the platform's own edit for an external source change.
                await refresh_registration_source_context_after_writeback(
                    cur,
                    parser_type=parser_type,
                    source_id=int(source["id"]),
                    previous_revision=int(source["revision"]),
                    previous_row_hash=str(source["row_hash"] or ""),
                    current_revision=int(revision),
                    current_row_hash=local_row_hash(after),
                )
        elif action == "cancel":
            await cancel_registration_link(
                cur,
                parser_type=parser_type,
                row_key=row_key_before,
                source_id=int(source["id"]),
                user_id=int(user.get("id")) if user.get("id") else None,
            )

    return transaction_prepare, transaction_callback, True


def _require_task_edit_user(user: dict) -> dict:
    if not (
        has_permission(user, ONLINE_RAW_EDIT)
        or has_permission(user, ONLINE_TASK_MANAGE)
    ):
        raise HTTPException(403, "当前权限组不能修改流口任务")
    if has_permission(user, ONLINE_RAW_EDIT):
        return user
    # online.task.manage 只在流口任务路由内借用现有逐行回写校验。
    # 不修改登录态中的权限，因此不会开放 Univer 或腾讯原始行管理。
    scoped_user = dict(user)
    scoped_user["permissions"] = list(dict.fromkeys([
        *(user.get("permissions") or []),
        ONLINE_RAW_EDIT,
    ]))
    return scoped_user


def _task_capability_user(user: dict) -> dict:
    if has_permission(user, ONLINE_TASK_MANAGE) and not has_permission(
        user, ONLINE_RAW_EDIT
    ):
        return _require_task_edit_user(user)
    return user


def _require_analysis_user(user: dict) -> dict:
    if not has_permission(user, ONLINE_TASK_MANAGE):
        raise HTTPException(403, "当前权限组不能处理网格核查研判")
    return _require_task_edit_user(user)


def _require_unverifiable_reviewer(user: dict) -> dict:
    position = str((user.get("member") or {}).get("position") or "").strip()
    if not (is_flow_task_admin(user) or position in {"基础管控", "中队长", "所队领导"}):
        raise HTTPException(403, "只有基础管控及以上岗位可以处理两级研判")
    return _require_task_edit_user(user)


def _can_assign_tasks(context: dict) -> bool:
    return bool(context.get("admin_mode")) or context.get("position") == "组长"


def _json_field(field: str) -> str:
    safe = field.replace('"', '\\"')
    return (
        "TRIM(COALESCE(JSON_UNQUOTE(JSON_EXTRACT(" 
        f"projection.values_json, '$.\"{safe}\"')), ''))"
    )


def _active_source_count_sql(parser_type: str, alias: str = "projection") -> str:
    return (
        "(SELECT COUNT(*) FROM _online_source_rows AS active_source "
        f"WHERE active_source.parser_type={alias}.parser_type "
        f"AND active_source.row_key={alias}.row_key "
        "AND active_source.archived_at IS NULL "
        f"{active_source_sql_filter(parser_type, 'active_source')})"
    )


def _source_exception_condition(parser_type: str) -> str:
    if local_data_source_enabled():
        return f"({_active_source_count_sql(parser_type)} > 1)"
    return "(projection.conflict=1 OR projection.source_count>1)"


def _review_condition(parser_type: str) -> str:
    workflow = TASK_WORKFLOWS[parser_type]
    conditions = [_source_exception_condition(parser_type)]
    if not workflow.valid_results:
        conditions.append(
            f"({_json_field(workflow.result_field)} LIKE '%%无法核实%%')"
        )
    if is_registration_task(parser_type):
        conditions.append(_registration_review_condition(parser_type))
    return "(" + " OR ".join(conditions) + ")"


def _registration_review_condition(parser_type: str) -> str:
    if not is_registration_task(parser_type):
        return "1=0"
    return (
        "EXISTS (SELECT 1 FROM _task_registration_links AS registration_link "
        "WHERE registration_link.parser_type=projection.parser_type "
        "AND registration_link.row_key=projection.row_key "
        "AND registration_link.status='review_required')"
    )


def _review_stage_condition(parser_type: str, stage: ReviewStage) -> str:
    workflow = TASK_WORKFLOWS[parser_type]
    if stage == "all":
        return "1=1"
    if supports_unverifiable_review(parser_type):
        structured = {
            "waiting_analysis": (INITIAL_PENDING,),
            "analyzed": (INITIAL_EXTENSION, DEEP_PENDING, DEEP_EXTENSION),
            "initial_pending": (INITIAL_PENDING,),
            "initial_extension": (INITIAL_EXTENSION,),
            "deep_pending": (DEEP_PENDING,),
            "deep_extension": (DEEP_EXTENSION,),
        }.get(stage)
        if structured:
            state_sql = ",".join("'" + value.replace("'", "''") + "'" for value in structured)
            # 保留旧的文本判定作为无流程历史数据的兼容回退。
            analysis = " OR ".join(
                f"{_json_field(field)}<>''" for field in workflow.analysis_fields
            ) or "0"
            unable = f"{_json_field(workflow.result_field)} LIKE '%%无法核实%%'"
            legacy = (
                f"({unable} AND NOT ({analysis}))"
                if stage in {"waiting_analysis", "initial_pending"}
                else f"({unable} AND ({analysis}))"
            )
            return (
                f"(EXISTS (SELECT 1 FROM _unverifiable_review_flows flow "
                f"WHERE flow.parser_type=projection.parser_type "
                f"AND flow.row_key=projection.row_key AND flow.state IN ({state_sql})) "
                f"OR ({legacy} AND NOT EXISTS (SELECT 1 FROM _unverifiable_review_flows flow2 "
                "WHERE flow2.parser_type=projection.parser_type AND flow2.row_key=projection.row_key)))"
            )
    if workflow.valid_results:
        return "1=0"
    analysis = " OR ".join(
        f"{_json_field(field)}<>''" for field in workflow.analysis_fields
    ) or "0"
    unable = f"{_json_field(workflow.result_field)} LIKE '%%无法核实%%'"
    return (
        f"({unable} AND ({analysis}))"
        if stage == "analyzed"
        else f"({unable} AND NOT ({analysis}))"
    )


def _priority_case(parser_type: str) -> str:
    workflow = TASK_WORKFLOWS[parser_type]
    source_exception = _source_exception_condition(parser_type)
    if workflow.valid_results:
        analyzed = "0"
        waiting = "0"
    else:
        analysis = " OR ".join(
            f"{_json_field(field)}<>''" for field in workflow.analysis_fields
        ) or "0"
        unable = f"{_json_field(workflow.result_field)} LIKE '%%无法核实%%'"
        analyzed = f"({unable} AND ({analysis}))"
        waiting = f"({unable} AND NOT ({analysis}))"
    return (
        "CASE "
        "WHEN projection.task_state='completed' THEN 'completed' "
        f"WHEN {analyzed} THEN 'analyzed' "
        f"WHEN {source_exception} THEN 'source_exception' "
        "WHEN projection.pending_state IN ('pending','retry','conflict') THEN 'pending_sync' "
        f"WHEN {waiting} THEN 'waiting_analysis' "
        "ELSE 'ordinary' END"
    )


def _priority_order(parser_type: str) -> str:
    bucket = _priority_case(parser_type)
    return (
        f"CASE {bucket} "
        "WHEN 'analyzed' THEN 0 "
        "WHEN 'source_exception' THEN 1 "
        "WHEN 'pending_sync' THEN 2 "
        "WHEN 'ordinary' THEN 3 "
        "WHEN 'waiting_analysis' THEN 4 "
        "ELSE 5 END"
    )


def _field_order(fields: tuple[str, ...]) -> str:
    candidates = ", ".join(f"NULLIF({_json_field(field)}, '')" for field in fields)
    raw_value = f"COALESCE({candidates}, '')" if candidates else "''"
    normalized = f"LOWER(REGEXP_REPLACE({raw_value}, '[[:space:]]+', ''))"
    return f"CASE WHEN {normalized}='' THEN 1 ELSE 0 END, {normalized}"


def _address_order(parser_type: str) -> str:
    return _field_order(TASK_WORKFLOWS[parser_type].address_fields)


def _identity_order(parser_type: str) -> str:
    return _field_order(TASK_WORKFLOWS[parser_type].identity_fields)


def _original_address_fields(parser_type: str) -> tuple[str, ...]:
    fields = tuple(
        field
        for field in TASK_WORKFLOWS[parser_type].address_fields
        if field != "现住址"
    )
    return fields or TASK_WORKFLOWS[parser_type].address_fields


def _original_address_order(parser_type: str) -> str:
    return _field_order(_original_address_fields(parser_type))


def _analysis_field_order(parser_types: list[str], field_name: str) -> str:
    field_cases: list[str] = []
    for parser_type in parser_types:
        workflow = TASK_WORKFLOWS[parser_type]
        fields = (
            _original_address_fields(parser_type)
            if field_name == "original_address_fields"
            else getattr(workflow, field_name)
        )
        candidates = ", ".join(
            f"NULLIF({_json_field(field)}, '')" for field in fields
        )
        value = f"COALESCE({candidates}, '')" if candidates else "''"
        safe_parser_type = parser_type.replace("'", "''")
        field_cases.append(
            f"WHEN projection.parser_type='{safe_parser_type}' THEN {value}"
        )
    raw_value = "CASE " + " ".join(field_cases) + " ELSE '' END"
    normalized = f"LOWER(REGEXP_REPLACE(({raw_value}), '[[:space:]]+', ''))"
    return f"CASE WHEN {normalized}='' THEN 1 ELSE 0 END, {normalized}"


def _assignment_candidate(
    parser_type: str,
    row_key: str,
    community: str,
    values: dict,
    small_community_name: str = "",
    match_status: str = "unmatched",
) -> dict[str, str]:
    summary = TASK_WORKFLOWS[parser_type].summary(values)
    return {
        "row_key": row_key,
        "community": community,
        "small_community": small_community_name,
        "match_status": match_status,
        "source": summary["source"] or TASK_WORKFLOWS[parser_type].label,
        "address": summary["address"] or "未填写地址",
    }


def _qmf_feedback_state_sql() -> str:
    raw_result = _json_field("核查结果")
    current_result = (
        f"CASE {raw_result} "
        "WHEN '离吴' THEN '离开不返吴' "
        "WHEN '近期反吴' THEN '近期返吴' "
        f"ELSE {raw_result} END"
    )
    return f"""
        COALESCE((
            SELECT CASE
                WHEN snapshot.error_code<>'' THEN 'error'
                WHEN snapshot.platform_result<>{current_result}
                  OR snapshot.last_scanned_at<DATE_SUB(UTC_TIMESTAMP(), INTERVAL 7 DAY)
                  OR source.id IS NULL
                  OR source.revision<>snapshot.source_revision
                  OR source.row_hash<>snapshot.source_row_hash
                THEN 'stale'
                WHEN snapshot.feedback_state='pending' THEN 'pending'
                WHEN snapshot.feedback_state='completed_match' THEN 'completed_match'
                WHEN snapshot.feedback_state='completed_mismatch' THEN 'completed_mismatch'
                WHEN snapshot.feedback_state='not_found' THEN 'not_found'
                WHEN snapshot.feedback_state='non_jurisdiction' THEN 'non_jurisdiction'
                ELSE 'error'
            END
            FROM _qmf_status_snapshots AS snapshot
            LEFT JOIN _online_source_rows AS source
              ON source.id=snapshot.source_id
             AND source.parser_type=snapshot.parser_type
             AND source.row_key=snapshot.row_key
            WHERE snapshot.parser_type=projection.parser_type
              AND snapshot.row_key=projection.row_key
            LIMIT 1
        ), 'not_scanned')
    """.strip()


async def _qmf_status_by_rows(
    cur,
    parser_type: str,
    rows: list[tuple],
) -> dict[str, dict]:
    if parser_type != MODEL_THREE_PARSER or not rows:
        return {}
    row_values = {
        str(row[0]): json_value(row[1], {})
        for row in rows
    }
    keys = list(row_values)
    placeholders = ",".join(["%s"] * len(keys))
    await cur.execute(
        f"""
        SELECT snapshot.row_key,snapshot.source_id,snapshot.source_revision,
               snapshot.source_row_hash,snapshot.platform_result,
               snapshot.feedback_state,snapshot.feedback_result,
               snapshot.checked_at,snapshot.origin,snapshot.error_code,
               snapshot.last_scanned_at,
               source.id,source.revision,source.row_hash
        FROM _qmf_status_snapshots AS snapshot
        LEFT JOIN _online_source_rows AS source
          ON source.id=snapshot.source_id
         AND source.parser_type=snapshot.parser_type
         AND source.row_key=snapshot.row_key
        WHERE snapshot.parser_type=%s
          AND snapshot.row_key IN ({placeholders})
        """,
        (parser_type, *keys),
    )
    result: dict[str, dict] = {}
    now = datetime.utcnow()
    for row in await cur.fetchall():
        row_key = str(row[0])
        platform_result = normalize_qmf_status_result(
            row_values.get(row_key, {}).get("核查结果")
        )
        stale = (
            str(row[4] or "") != platform_result
            or row[11] is None
            or int(row[12] or 0) != int(row[2] or 0)
            or str(row[13] or "") != str(row[3] or "")
            or not row[10]
            or (now - row[10]).total_seconds() > 7 * 24 * 60 * 60
        )
        raw_state = str(row[5] or "")
        state = (
            "error" if str(row[9] or "")
            else "stale" if stale
            else raw_state if raw_state in {
                "pending", "completed_match", "completed_mismatch", "not_found"
            }
            else "error"
        )
        result[row_key] = {
            "state": state,
            "platform_result": str(row[4] or ""),
            "feedback_result": str(row[6] or ""),
            "checked_at": str(row[7] or ""),
            "origin": str(row[8] or ""),
            "error_code": str(row[9] or ""),
            "last_scanned_at": _iso_utc(row[10]),
        }
    for row_key, values in row_values.items():
        result.setdefault(row_key, {
            "state": "not_scanned",
            "platform_result": normalize_qmf_status_result(values.get("核查结果")),
            "feedback_result": "",
            "checked_at": "",
            "origin": "",
            "error_code": "",
            "last_scanned_at": None,
        })
    return result


def _priority_bucket(
    parser_type: str,
    values: dict,
    source_count: int,
    conflict: bool,
    pending: bool,
    task_state_value: str,
) -> str:
    workflow = TASK_WORKFLOWS[parser_type]
    if task_state_value == "completed" or workflow.state(values) == "completed":
        return "completed"
    if workflow.review_stage(values) == "analyzed":
        return "analyzed"
    if conflict or source_count > 1:
        return "source_exception"
    if pending:
        return "pending_sync"
    if workflow.review_stage(values) == "waiting_analysis":
        return "waiting_analysis"
    return "ordinary"


def _multi_filter_condition(
    column: str,
    values: list[str],
) -> tuple[str, list[str]]:
    normalized = list(dict.fromkeys(str(value).strip() for value in values if str(value).strip()))
    if not normalized:
        return "1=1", []
    include_empty = EMPTY_FILTER_VALUE in normalized
    non_empty = [value for value in normalized if value != EMPTY_FILTER_VALUE]
    predicates: list[str] = []
    params: list[str] = []
    if non_empty:
        placeholders = ", ".join(["%s"] * len(non_empty))
        predicates.append(f"projection.{column} IN ({placeholders})")
        params.extend(non_empty)
    if include_empty:
        predicates.append(f"TRIM(COALESCE(projection.{column}, ''))='' ")
    return "(" + " OR ".join(predicates) + ")", params


def _task_where(
    context: dict,
    parser_type: str,
    data: TaskSearch,
    *,
    include_priority: bool = True,
) -> tuple[str, list]:
    scope_where, scope_params = _scope_where(context, data.scope)
    where_parts = ["projection.parser_type=%s", scope_where]
    params: list = [parser_type, *scope_params]
    if local_data_source_enabled():
        where_parts.append(
            "EXISTS (SELECT 1 FROM _online_source_rows AS active_source "
            "WHERE active_source.parser_type=projection.parser_type "
            "AND active_source.row_key=projection.row_key "
            "AND active_source.archived_at IS NULL "
            f"{active_source_sql_filter(parser_type, 'active_source')})"
        )
    review_condition = _review_condition(parser_type)
    if data.status == "pending":
        where_parts.append("projection.task_state<>'completed'")
    elif data.status == "unchecked":
        where_parts.append("projection.task_state='unchecked'")
    elif data.status == "checked":
        where_parts.append("projection.task_state='checked'")
    elif data.status == "completed":
        where_parts.append("projection.task_state='completed'")
    elif data.status == "review":
        where_parts.append(review_condition)
    elif data.status == "registration_review":
        where_parts.append(_registration_review_condition(parser_type))
    if data.review_stage != "all":
        where_parts.append(_review_stage_condition(parser_type, data.review_stage))
    community_condition, community_params = _multi_filter_condition(
        "community", data.communities
    )
    small_community_condition, small_community_params = _multi_filter_condition(
        "small_community_name", data.small_communities
    )
    match_status_condition, match_status_params = _multi_filter_condition(
        "address_match_status", data.match_status
    )
    inspector_condition, inspector_params = _multi_filter_condition(
        "inspector", data.inspectors
    )
    where_parts.extend([
        community_condition, small_community_condition,
        match_status_condition, inspector_condition,
    ])
    params.extend(community_params)
    params.extend(small_community_params)
    params.extend(match_status_params)
    params.extend(inspector_params)
    keyword = data.keyword.strip()
    if keyword:
        where_parts.append("projection.search_text LIKE %s")
        params.append(f"%{keyword}%")
    if data.watch_categories:
        if not settings.REGISTRY_FEATURE_ENABLED:
            where_parts.append("1=0")
        else:
            placeholders = ",".join(["%s"] * len(data.watch_categories))
            registry = settings.MYSQL_REGISTRY_DB.replace("`", "")
            where_parts.append(
                f"EXISTS (SELECT 1 FROM `{registry}`.online_task_watch_snapshots watch_snapshot "
                f"JOIN `{registry}`.watch_assignments watch_assignment "
                "ON watch_assignment.id=watch_snapshot.assignment_id "
                "WHERE watch_snapshot.parser_type=projection.parser_type "
                "AND watch_snapshot.row_key=projection.row_key "
                f"AND watch_assignment.category_id IN ({placeholders}))"
            )
            params.extend(data.watch_categories)
    if data.qmf_feedback_states:
        if parser_type != MODEL_THREE_PARSER:
            where_parts.append("1=0")
        else:
            placeholders = ",".join(["%s"] * len(data.qmf_feedback_states))
            where_parts.append(
                f"({_qmf_feedback_state_sql()}) IN ({placeholders})"
            )
            params.extend(data.qmf_feedback_states)
    if include_priority and data.priority != "all":
        where_parts.append(f"({_priority_case(parser_type)})=%s")
        params.append(data.priority)
    return " AND ".join(where_parts), params


def _analysis_stage_condition(
    parser_types: list[str],
    stage: ReviewStage,
) -> tuple[str, list[str]]:
    conditions: list[str] = []
    params: list[str] = []
    stages = (stage,) if stage != "all" else (
        "initial_pending", "initial_extension", "deep_pending", "deep_extension"
    )
    for parser_type in parser_types:
        stage_conditions = [
            _review_stage_condition(parser_type, current_stage)
            for current_stage in stages
        ]
        conditions.append(
            "(projection.parser_type=%s AND ("
            + " OR ".join(stage_conditions)
            + "))"
        )
        params.append(parser_type)
    return "(" + " OR ".join(conditions) + ")", params


def _analysis_task_where(
    context: dict,
    data: AnalysisTaskSearch,
    *,
    review_stage: ReviewStage | None = None,
) -> tuple[str, list]:
    parser_types = list(dict.fromkeys(data.parser_types))
    type_placeholders = ",".join(["%s"] * len(parser_types))
    scope_where, scope_params = _scope_where(context, "all")
    where_parts = [
        f"projection.parser_type IN ({type_placeholders})",
        scope_where,
    ]
    params: list = [*parser_types, *scope_params]
    stage_sql, stage_params = _analysis_stage_condition(
        parser_types,
        review_stage if review_stage is not None else data.review_stage,
    )
    where_parts.append(stage_sql)
    params.extend(stage_params)
    community_condition, community_params = _multi_filter_condition(
        "community", data.communities
    )
    small_community_condition, small_community_params = _multi_filter_condition(
        "small_community_name", data.small_communities
    )
    match_status_condition, match_status_params = _multi_filter_condition(
        "address_match_status", data.match_status
    )
    inspector_condition, inspector_params = _multi_filter_condition(
        "inspector", data.inspectors
    )
    where_parts.extend([
        community_condition, small_community_condition,
        match_status_condition, inspector_condition,
    ])
    params.extend(community_params)
    params.extend(small_community_params)
    params.extend(match_status_params)
    params.extend(inspector_params)
    if data.keyword.strip():
        where_parts.append("projection.search_text LIKE %s")
        params.append(f"%{data.keyword.strip()}%")
    if data.watch_categories:
        if not settings.REGISTRY_FEATURE_ENABLED:
            where_parts.append("1=0")
        else:
            placeholders = ",".join(["%s"] * len(data.watch_categories))
            registry = settings.MYSQL_REGISTRY_DB.replace("`", "")
            where_parts.append(
                f"EXISTS (SELECT 1 FROM `{registry}`.online_task_watch_snapshots watch_snapshot "
                f"JOIN `{registry}`.watch_assignments watch_assignment "
                "ON watch_assignment.id=watch_snapshot.assignment_id "
                "WHERE watch_snapshot.parser_type=projection.parser_type "
                "AND watch_snapshot.row_key=projection.row_key "
                f"AND watch_assignment.category_id IN ({placeholders}))"
            )
            params.extend(data.watch_categories)
    return " AND ".join(where_parts), params


def _analysis_order(data: AnalysisTaskSearch) -> str:
    parser_types = list(dict.fromkeys(data.parser_types))
    stage_predicates: dict[str, list[str]] = {
        "initial_pending": [],
        "initial_extension": [],
        "deep_pending": [],
        "deep_extension": [],
    }
    for parser_type in parser_types:
        for stage in stage_predicates:
            stage_predicates[stage].append(
                f"(projection.parser_type='{parser_type}' AND "
                f"{_review_stage_condition(parser_type, stage)})"
            )
    stage_order = "CASE "
    for index, stage in enumerate(stage_predicates):
        predicates = stage_predicates[stage]
        if predicates:
            stage_order += f"WHEN {' OR '.join(predicates)} THEN {index} "
    stage_order += "ELSE 4 END"
    if data.sort == "address_asc":
        return (
            f"{stage_order}, "
            f"{_analysis_field_order(parser_types, 'original_address_fields')}, "
            "projection.row_key"
        )
    if data.sort == "identity_asc":
        return (
            f"{stage_order}, "
            f"{_analysis_field_order(parser_types, 'identity_fields')}, "
            f"{_analysis_field_order(parser_types, 'original_address_fields')}, "
            "projection.row_key"
        )
    if data.sort == "updated_asc":
        return f"{stage_order}, projection.updated_at ASC, projection.row_key"
    if data.sort == "updated_desc":
        return f"{stage_order}, projection.updated_at DESC, projection.row_key"
    return f"{stage_order}, projection.updated_at DESC, projection.row_key"


def _task_order(parser_type: str, sort: SortMode) -> str:
    completed_last = "CASE WHEN projection.task_state='completed' THEN 1 ELSE 0 END"
    if sort == "updated_asc":
        return f"{completed_last}, projection.updated_at ASC, projection.row_key"
    if sort == "updated_desc":
        return f"{completed_last}, projection.updated_at DESC, projection.row_key"
    if sort == "address_asc":
        return f"{_original_address_order(parser_type)}, projection.row_key"
    if sort == "identity_asc":
        return (
            f"{_identity_order(parser_type)}, "
            f"{_original_address_order(parser_type)}, projection.row_key"
        )
    return (
        f"{_priority_order(parser_type)}, {_address_order(parser_type)}, "
        "projection.updated_at DESC, projection.row_key"
    )


async def _analysis_filter_options(
    cur,
    context: dict,
    user: dict,
    data: AnalysisTaskSearch,
) -> dict:
    option_data = data.model_copy(update={
        "communities": [],
    })
    base_where, base_params = _analysis_task_where(
        context,
        option_data,
        review_stage=data.review_stage,
    )
    community_condition, community_params = _multi_filter_condition(
        "community", data.communities
    )
    inspector_where = f"{base_where} AND {community_condition}"
    inspector_params = [*base_params, *community_params]
    result = {"communities": [], "inspectors": [], "watch_categories": []}
    for column, key, empty_label, option_where, option_params in (
        ("community", "communities", "社区未填写", base_where, base_params),
        ("inspector", "inspectors", "未分配核查人", inspector_where, inspector_params),
    ):
        await cur.execute(
            f"""
            SELECT projection.{column}, COUNT(*)
            FROM _online_source_projection AS projection
            WHERE {option_where}
            GROUP BY projection.{column}
            ORDER BY CASE WHEN TRIM(COALESCE(projection.{column}, ''))='' THEN 1 ELSE 0 END,
                     projection.{column}
            """,
            option_params,
        )
        for value, count in await cur.fetchall():
            normalized = str(value or "").strip()
            result[key].append({
                "value": normalized or EMPTY_FILTER_VALUE,
                "label": normalized or empty_label,
                "count": int(count or 0),
            })
    if settings.REGISTRY_FEATURE_ENABLED:
        registry = settings.MYSQL_REGISTRY_DB.replace("`", "")
        await cur.execute(
            f"""
            SELECT category.id, category.name, category.color, category.alert_level,
                   COUNT(DISTINCT projection.row_key)
            FROM _online_source_projection projection
            JOIN `{registry}`.online_task_watch_snapshots snapshot
              ON snapshot.parser_type=projection.parser_type
             AND snapshot.row_key=projection.row_key
            JOIN `{registry}`.watch_assignments assignment
              ON assignment.id=snapshot.assignment_id
            JOIN `{registry}`.watch_categories category
              ON category.id=assignment.category_id
            WHERE {base_where}
            GROUP BY category.id, category.name, category.color, category.alert_level
            ORDER BY category.sort_order, category.id
            """,
            base_params,
        )
        result["watch_categories"] = [
            {"value": int(row[0]), "label": str(row[1]), "color": str(row[2]),
             "alert_level": str(row[3]), "count": int(row[4] or 0)}
            for row in await cur.fetchall()
        ]
    assignment_context = await inspector_option_context(
        cur,
        _task_capability_user(user),
        assignment_only=True,
    )
    result["assignment"] = {
        "enabled": False,
        "community_aliases": assignment_context["community_aliases"],
        "inspectors_by_community": assignment_context["inspectors_by_community"],
    }
    return result


async def _list_analysis_tasks_data(
    data: AnalysisTaskSearch,
    user: dict,
    conn,
) -> dict:
    parser_types = list(dict.fromkeys(data.parser_types))
    if not parser_types or any(parser_type not in UNVERIFIABLE_REVIEW_TYPES for parser_type in parser_types):
        raise HTTPException(400, "存在尚未接入研判工作台的业务表")
    context = await _flow_context(conn, user)
    where_sql, query_params = _analysis_task_where(context, data)
    base_data = data.model_copy(update={"review_stage": "all"})
    base_where, base_params = _analysis_task_where(
        context,
        base_data,
        review_stage="all",
    )
    async with conn.cursor() as cur:
        ready_values = []
        for parser_type in parser_types:
            ready_values.append(await _task_source_ready(cur, parser_type))
        if not all(ready_values):
            return {
                "data": [], "total": 0, "page": data.page, "page_size": data.page_size,
                "source_ready": False,
                "message": (
                    "部分业务表本地来源尚未建立"
                    if local_data_source_enabled()
                    else "部分业务表来源尚未建立，请等待一次正常同步"
                ),
                "facets": _empty_facets(), "priority_labels": PRIORITY_LABELS,
                "filters": {"parser_types": parser_types, "scope": data.scope,
                    "review_stage": data.review_stage, "communities": data.communities,
                    "inspectors": data.inspectors, "watch_categories": data.watch_categories,
                    "sort": data.sort, "keyword_present": bool(data.keyword.strip())},
            }
        await cur.execute(
            f"SELECT COUNT(*) FROM _online_source_projection AS projection WHERE {where_sql}",
            query_params,
        )
        total = int((await cur.fetchone())[0] or 0)
        facets = _empty_facets()
        await cur.execute(
            f"SELECT COUNT(*) FROM _online_source_projection AS projection WHERE {base_where}",
            base_params,
        )
        facets["total"] = int((await cur.fetchone())[0] or 0)
        for stage, key in (
            ("initial_pending", "initial_pending"),
            ("initial_extension", "initial_extension"),
            ("deep_pending", "deep_pending"),
            ("deep_extension", "deep_extension"),
        ):
            stage_where, stage_params = _analysis_task_where(context, base_data, review_stage=stage)
            await cur.execute(
                f"SELECT COUNT(*) FROM _online_source_projection AS projection WHERE {stage_where}",
                stage_params,
            )
            facets["review_stage_counts"][key] = int((await cur.fetchone())[0] or 0)
        await cur.execute(
            f"""
            SELECT projection.task_state, COUNT(*)
            FROM _online_source_projection AS projection
            WHERE {base_where}
            GROUP BY projection.task_state
            """,
            base_params,
        )
        for state, count in await cur.fetchall():
            if str(state) in facets["status_counts"]:
                facets["status_counts"][str(state)] = int(count or 0)
        await cur.execute(
            f"""
            SELECT projection.parser_type, projection.row_key, projection.values_json,
                   projection.source_count, projection.conflict,
                   projection.pending_state, projection.task_state
            FROM _online_source_projection AS projection
            WHERE {where_sql}
            ORDER BY {_analysis_order(data)}
            LIMIT %s OFFSET %s
            """,
            [*query_params, data.page_size, (data.page - 1) * data.page_size],
        )
        rows = await cur.fetchall()
        review_by_row = await review_flows_by_rows(
            cur, [(str(row[0]), str(row[1])) for row in rows]
        )
        watch_by_parser: dict[str, dict[str, dict]] = {}
        for parser_type in parser_types:
            keys = [str(row[1]) for row in rows if str(row[0]) == parser_type]
            watch_by_parser[parser_type] = (
                await task_watch_payload(cur, parser_type, keys)
                if settings.REGISTRY_FEATURE_ENABLED and keys else {}
            )
        address_matches_by_parser = {}
        for current_type in parser_types:
            keys = [str(row[1]) for row in rows if str(row[0]) == current_type]
            address_matches_by_parser[current_type] = await _address_matches_by_rows(
                cur, current_type, keys
            )
    photo_fetched: dict[str, set[str]] = {}
    for parser_type in parser_types:
        keys = [str(row[1]) for row in rows if str(row[0]) == parser_type]
        photo_fetched[parser_type] = await _task_photo_fetched_rows(user, parser_type, keys)
    return {
        "data": [
            _task_record(
                str(row[0]), str(row[1]), json_value(row[2], {}), int(row[3] or 0),
                bool(row[4]), bool(str(row[5] or "")), str(row[6] or ""),
                watch_by_parser.get(str(row[0]), {}).get(str(row[1])),
                str(row[1]) in photo_fetched.get(str(row[0]), set()),
                str(row[5] or ""),
                review_flow=review_by_row.get((str(row[0]), str(row[1]))),
                address_match=address_matches_by_parser.get(str(row[0]), {}).get(str(row[1])),
            )
            for row in rows
        ],
        "total": total,
        "page": data.page,
        "page_size": data.page_size,
        "source_ready": True,
        "message": "",
        "facets": facets,
        "priority_labels": PRIORITY_LABELS,
        "filters": {"parser_types": parser_types, "scope": data.scope,
            "review_stage": data.review_stage, "communities": data.communities,
            "inspectors": data.inspectors, "watch_categories": data.watch_categories,
            "sort": data.sort, "keyword_present": bool(data.keyword.strip())},
    }


async def _flow_context(conn, user: dict) -> dict:
    if is_flow_task_elevated(user):
        # online.task.manage 是流口任务工作台内的全所管理能力。
        # 它不会改变用户在 Univer、辖区档案等其他模块的数据范围。
        view_communities = (
            None
            if has_permission(user, ONLINE_TASK_MANAGE)
            else effective_view_communities(user)
        )
        community_values = (
            None
            if view_communities is None
            else await community_names_for_scopes(conn, view_communities)
        )
        member = user.get("member") or {}
        return {
            "name": str(
                member.get("name")
                or user.get("display_name")
                or user.get("username")
                or "管理员"
            ).strip(),
            "position": str(member.get("position") or "管理员").strip(),
            "community": "全所",
            "community_values": community_values,
            "admin_mode": True,
        }
    name, community = require_flow_user(user)
    aliases = await community_names_for_scopes(conn, [community])
    return {
        "name": name,
        "position": str((user.get("member") or {}).get("position") or ""),
        "community": community,
        "community_values": aliases or [community],
        "admin_mode": False,
    }


def _scope_where(
    context: dict,
    scope: FlowScope,
    *,
    alias: str = "projection",
) -> tuple[str, list[str]]:
    if scope == "all" and not context.get("admin_mode"):
        raise HTTPException(403, "组员和组长只能查看本人或本社区任务")
    community_values = context["community_values"]
    if community_values is None:
        clause = "1=1"
        params: list[str] = []
    elif community_values:
        placeholders = ", ".join(["%s"] * len(community_values))
        clause = f"{alias}.community IN ({placeholders})"
        params = list(community_values)
    else:
        clause = "1=0"
        params = []
    if scope == "mine":
        clause += f" AND LOWER(TRIM({alias}.inspector))=LOWER(TRIM(%s))"
        params.append(context["name"])
    return clause, params


async def _source_readiness(cur) -> dict[str, bool]:
    result: dict[str, bool] = {}
    for parser_type in MOBILE_TASK_TYPES:
        result[parser_type] = await _task_source_ready(cur, parser_type)
    return result


async def _aggregate_live(
    cur,
    context: dict,
    scope: FlowScope,
    parser_types: tuple[str, ...] = MOBILE_TASK_TYPES,
) -> dict[str, dict]:
    where, params = _scope_where(context, scope)
    type_placeholders = ", ".join(["%s"] * len(parser_types))
    active_source_exists = "1=1"
    source_exception = "projection.conflict=1 OR projection.source_count>1"
    if local_data_source_enabled():
        active_source_exists = (
            "EXISTS (SELECT 1 FROM _online_source_rows AS active_source "
            "WHERE active_source.parser_type=projection.parser_type "
            "AND active_source.row_key=projection.row_key "
            "AND active_source.archived_at IS NULL "
            f"{active_source_sql_filter('all', 'active_source')})"
        )
        source_exception = (
            "(SELECT COUNT(*) FROM _online_source_rows AS active_source "
            "WHERE active_source.parser_type=projection.parser_type "
            "AND active_source.row_key=projection.row_key "
            "AND active_source.archived_at IS NULL "
            f"{active_source_sql_filter('all', 'active_source')}) > 1"
        )
    await cur.execute(
        f"""
        SELECT projection.parser_type, projection.task_state,
               COUNT(*),
               SUM(CASE WHEN {source_exception}
                              OR EXISTS (
                                  SELECT 1 FROM _task_registration_links registration_link
                                  WHERE registration_link.parser_type=projection.parser_type
                                    AND registration_link.row_key=projection.row_key
                                    AND registration_link.status='review_required'
                              )
                              OR (projection.parser_type IN (
                                      '全链条','出租房屋核查','寄递业'
                                  ) AND {_json_field('核查结果')}
                                      LIKE '%%无法核实%%')
                              OR (projection.parser_type='疑似返苏'
                                  AND {_json_field('核查反馈')}
                                      LIKE '%%无法核实%%')
                        THEN 1 ELSE 0 END)
        FROM _online_source_projection AS projection
        WHERE projection.parser_type IN ({type_placeholders})
          AND {where}
          AND {active_source_exists}
        GROUP BY projection.parser_type, projection.task_state
        """,
        (*parser_types, *params),
    )
    totals = {
        parser_type: {
            "parser_type": parser_type,
            "label": TASK_WORKFLOWS[parser_type].label,
            "pending": 0,
            "unchecked": 0,
            "checked": 0,
            "completed": 0,
            "review": 0,
        }
        for parser_type in parser_types
    }
    for parser_type, state, count, review in await cur.fetchall():
        item = totals[str(parser_type)]
        normalized_state = str(state or "unchecked")
        amount = int(count or 0)
        if normalized_state in {"unchecked", "checked", "completed"}:
            item[normalized_state] += amount
        if normalized_state != "completed":
            item["pending"] += amount
        item["review"] += int(review or 0)
    return totals


async def _validate_assignment(
    cur,
    context: dict,
    changes: dict[str, str],
) -> None:
    if "核查人" not in changes:
        return
    if context.get("admin_mode"):
        return
    if context["position"] != "组长":
        raise HTTPException(403, "只有组长可以在手机任务中转派任务")
    assignee = str(changes.get("核查人") or "").strip()
    if not assignee:
        raise HTTPException(400, "请选择要分配的组员")
    await cur.execute(
        """
        SELECT member.id
        FROM _grid_members AS member
        JOIN _grid_member_department_links AS link
          ON link.member_id=member.id
        JOIN _departments AS department
          ON department.id=link.department_id
         AND department.department_type='community'
         AND department.is_active=1
        JOIN _communities AS community
          ON community.id=department.community_id
        WHERE community.name=%s
          AND member.name=%s
          AND member.position IN ('组长', '组员')
          AND member.status='在岗'
        LIMIT 1
        """,
        (context["community"], assignee),
    )
    if not await cur.fetchone():
        raise HTTPException(400, "只能分配给本社区在岗组长或组员")


def _task_record(
    parser_type: str,
    row_key: str,
    values: dict,
    source_count: int,
    conflict: bool,
    pending: bool,
    task_state_value: str,
    watch: dict | None = None,
    photo_fetched: bool = False,
    sync_state: str = "",
    qmf_status: dict | None = None,
    residence_status: dict | None = None,
    registration_link: dict | None = None,
    review_flow: dict | None = None,
    address_match: dict | None = None,
) -> dict:
    workflow = TASK_WORKFLOWS[parser_type]
    normalized = {key: str(value or "") for key, value in values.items()}
    watch = watch or {}
    structured_stage = str((review_flow or {}).get("state") or "")
    # 正式结果已经完成后，旧的无法核实流程只作为历史记录保留，不能
    # 再把任务标成“流程已暂停”。来源异常仍可通过独立的来源/同步标记
    # 展示，不影响已完成状态的主口径。
    structured_active = (
        structured_stage in UNVERIFIABLE_ACTIVE_STATES
        and task_state_value != "completed"
    )
    return {
        "task_key": f"{parser_type}:{row_key}",
        "row_key": str(row_key),
        "parser_type": parser_type,
        "summary": workflow.summary(normalized),
        "community": str(normalized.get("社区") or normalized.get("下发社区") or ""),
        "inspector": str(normalized.get("核查人") or ""),
        "state": task_state_value or workflow.state(normalized),
        "needs_review": workflow.needs_review(
            normalized,
            source_count=int(source_count or 0),
            conflict=bool(conflict),
        ),
        "review_stage": structured_stage if structured_active else workflow.review_stage(normalized),
        "review_flow": review_flow,
        "source_count": int(source_count or 0),
        "conflict": bool(conflict),
        "pending_sync": bool(pending),
        "sync_state": sync_state or ("pending" if pending else ""),
        "photo_fetched": bool(photo_fetched),
        "priority": _priority_bucket(
            parser_type,
            normalized,
            int(source_count or 0),
            bool(conflict),
            bool(pending),
            task_state_value,
        ),
        "watch_marks": list(watch.get("watch_marks") or []),
        "first_dispatch_at": _iso_utc(watch.get("first_dispatch_at")),
        "qmf_status": qmf_status,
        "residence_status": residence_status,
        "registration_link": registration_link,
        "address_match": address_match,
    }


async def _address_matches_by_rows(
    cur,
    parser_type: str,
    row_keys: list[str],
) -> dict[str, dict]:
    if not row_keys:
        return {}
    placeholders = ",".join(["%s"] * len(row_keys))
    await cur.execute(
        f"""
        SELECT row_key, small_community_id, small_community_name,
               address_match_status, address_match_score,
               address_match_method, address_match_reason,
               address_match_candidates, address_match_version
        FROM _online_source_projection
        WHERE parser_type=%s AND row_key IN ({placeholders})
        """,
        (parser_type, *row_keys),
    )
    return {
        str(row[0]): {
            "small_community_id": int(row[1]) if row[1] is not None else None,
            "small_community_name": str(row[2] or ""),
            "status": str(row[3] or "unmatched"),
            "score": float(row[4] or 0),
            "method": str(row[5] or ""),
            "reason": str(row[6] or ""),
            "candidates": json_value(row[7], []),
            "version": str(row[8] or ""),
        }
        for row in await cur.fetchall()
    }


def _export_value(values: dict, *keys: str) -> str:
    for key in keys:
        value = values.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _mobile_export_row(
    parser_type: str,
    row_key: str,
    values: dict,
    source_id: int | None,
    source_count: int,
    conflict: bool,
    pending: bool,
    task_state: str,
    source_revision: int | None = None,
    source_row_hash: str = "",
    review_flow: dict | None = None,
    include_internal: bool = False,
) -> list[object]:
    workflow = TASK_WORKFLOWS[parser_type]
    stage = str((review_flow or {}).get("state") or workflow.review_stage(values) or "")
    public_values = [
        _export_value(values, "姓名", "核查对象", "对象姓名"),
        _export_value(values, "身份证号", "身份证号码", "公民身份号码"),
        _export_value(
            values,
            "联系号码",
            "手机号",
            "联系电话",
            "手机号码",
            "电话号码",
            "联系方式",
            "电话",
        ),
        _export_value(values, "原地址", "原住址", "地址", "疑似现住址"),
        _export_value(values, "现住址", "核查补充信息", "拟登记住址"),
        _export_value(values, "社区", "下发社区"),
        _export_value(values, "核查人"),
        task_state,
        UNVERIFIABLE_STATE_LABELS.get(stage, stage),
        _export_value(values, *workflow.analysis_fields),
        _export_value(values, *workflow.secondary_fields),
        _export_value(values, workflow.result_field),
        _export_value(values, *workflow.date_fields),
    ]
    if not include_internal:
        return public_values
    return [
        parser_type,
        row_key,
        source_id or "",
        source_revision or "",
        source_row_hash,
        int((review_flow or {}).get("flow_version") or 0) or "",
        *public_values,
        "",
        "是" if conflict or source_count != 1 else "否",
        "是" if pending else "否",
    ]


async def _mobile_export_workbook(
    *,
    data: TaskSearch | AnalysisTaskSearch,
    parser_type: str | None,
    user: dict,
    conn,
) -> tuple[BytesIO, int]:
    context = await _flow_context(conn, user)
    active_source_filter = (
        " AND source.spreadsheet_id=0 "
        "AND source.source_kind IN ('local_table','local_dispatch')"
        if local_data_source_enabled() else ""
    )
    if isinstance(data, AnalysisTaskSearch):
        parser_types = list(dict.fromkeys(data.parser_types))
        if not parser_types or any(value not in TASK_WORKFLOWS for value in parser_types):
            raise HTTPException(400, "存在尚未接入研判工作台的业务表")
        where_sql, query_params = _analysis_task_where(context, data)
        order_sql = _analysis_order(data)
        type_label = "研判任务"
        include_internal_columns = True
    else:
        if not parser_type or parser_type not in TASK_WORKFLOWS:
            raise HTTPException(400, "该业务尚未接入手机任务工作台")
        parser_types = [parser_type]
        where_sql, query_params = _task_where(context, parser_type, data)
        order_sql = _task_order(parser_type, data.sort)
        type_label = parser_type
        include_internal_columns = False
    async with conn.cursor() as cur:
        await cur.execute(
            f"""
            SELECT projection.parser_type, projection.row_key, projection.values_json,
                   projection.source_count, projection.conflict, projection.pending_state,
                   projection.task_state,
                   (SELECT MIN(source.id) FROM _online_source_rows source
                    WHERE source.parser_type=projection.parser_type
                      AND source.row_key=projection.row_key
                      AND source.archived_at IS NULL{active_source_filter}) AS source_id,
                   (SELECT source.revision FROM _online_source_rows source
                    WHERE source.parser_type=projection.parser_type
                      AND source.row_key=projection.row_key
                      AND source.archived_at IS NULL{active_source_filter}
                    ORDER BY source.id LIMIT 1) AS source_revision,
                   (SELECT source.row_hash FROM _online_source_rows source
                    WHERE source.parser_type=projection.parser_type
                      AND source.row_key=projection.row_key
                      AND source.archived_at IS NULL{active_source_filter}
                    ORDER BY source.id LIMIT 1) AS source_row_hash
            FROM _online_source_projection AS projection
            WHERE {where_sql}
            ORDER BY {order_sql}
            LIMIT 100000
            """,
            query_params,
        )
        rows = await cur.fetchall()
        review_by_row = await review_flows_by_rows(
            cur, [(str(row[0]), str(row[1])) for row in rows]
        )
    export_rows = [
        _mobile_export_row(
            str(row[0]), str(row[1]), json_value(row[2], {}),
            int(row[7]) if row[7] is not None else None,
            int(row[3] or 0), bool(row[4]), bool(row[5]), str(row[6] or ""),
            int(row[8]) if row[8] is not None else None,
            str(row[9] or ""),
            review_by_row.get((str(row[0]), str(row[1]))),
            include_internal=include_internal_columns,
        )
        for row in rows
    ]
    if include_internal_columns:
        headers = [
            "业务类型", "任务标识", "来源ID", "来源版本", "来源行哈希", "流程版本",
            "姓名", "身份证号", "手机号", "原地址", "现住址", "社区", "核查人",
            "任务状态", "研判阶段", "本次研判决定", "研判意见", "复核反馈", "核查结果", "截止日期",
            "来源异常", "待同步",
        ]
    else:
        headers = [
            "姓名", "身份证号", "手机号", "原地址", "现住址", "社区", "核查人",
            "任务状态", "研判阶段", "研判意见", "复核反馈", "核查结果", "截止日期",
        ]
    workbook = build_xlsx(
        type_label,
        headers,
        export_rows,
    )
    return workbook, len(export_rows)


def _source_in_community(
    parser,
    values: dict,
    community_values: list[str] | None,
) -> bool:
    if community_values is None:
        return True
    source_community = parser.community_value(values).strip()
    return bool(source_community and source_community in set(community_values))


@router.get("/home")
async def get_mobile_task_home(
    scope: FlowScope = Query("mine"),
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    context = await _flow_context(conn, user)
    async with conn.cursor() as cur:
        business_date = await get_business_date(cur)
        readiness = await _source_readiness(cur)
        selected = await _aggregate_live(cur, context, scope, MOBILE_TASK_TYPES)
        personal_scope: FlowScope = "all" if context["admin_mode"] else "mine"
        community_scope: FlowScope = "all" if context["admin_mode"] else "community"
        personal = await _aggregate_live(cur, context, personal_scope, SUMMARY_TASK_TYPES)
        community = await _aggregate_live(cur, context, community_scope, SUMMARY_TASK_TYPES)
        await cur.execute(
            "SELECT MAX(finished_at) FROM _sync_log "
            "WHERE status IN ('success', 'completed')"
        )
        sync_row = await cur.fetchone()

    date_text = business_date.isoformat()
    try:
        today = await get_online_overview(
            date_text,
            date_text,
            SUMMARY_TYPE,
            context["community_values"],
            inspector=None if context["admin_mode"] else context["name"],
            parser_types_override=list(SUMMARY_TASK_TYPES),
        )
    except (RuntimeError, ValueError):
        today = {"exists": False}

    personal_totals = list(personal.values())
    community_totals = list(community.values())
    businesses = sorted(
        selected.values(),
        key=lambda item: (item["pending"] == 0, -item["pending"], item["label"]),
    )
    for item in businesses:
        item["source_ready"] = bool(readiness.get(item["parser_type"]))
    return {
        "business_date": date_text,
        "last_success_at": _iso_utc(sync_row[0]) if sync_row else None,
        "scope": scope,
        "admin_mode": context["admin_mode"],
        "person": {
            "name": context["name"],
            "position": context["position"],
            "community": context["community"],
        },
        "personal": {
            "pending": sum(item["pending"] for item in personal_totals),
            "new_today": int(today.get("new_tasks") or 0) if today.get("exists") else None,
            "carryover_today": int(today.get("carryover_tasks") or 0) if today.get("exists") else None,
            "completed_today": int(today.get("completed_tasks") or 0) if today.get("exists") else None,
        },
        "community": {
            "pending": sum(item["pending"] for item in community_totals),
        },
        "daily_snapshot_available": bool(today.get("exists")),
        "businesses": businesses,
    }


def _empty_facets() -> dict:
    return {
        "total": 0,
        "priority_counts": {key: 0 for key in PRIORITY_KEYS},
        "review_stage_counts": {
            key: 0 for key in (
                "initial_pending", "initial_extension", "deep_pending", "deep_extension",
            )
        },
        "status_counts": {key: 0 for key in ("unchecked", "checked", "completed")},
        "registration_review_count": 0,
        "qmf_feedback_counts": {
            key: 0 for key in (
                "not_scanned", "stale", "pending", "completed_match",
                "completed_mismatch", "not_found", "non_jurisdiction", "error",
            )
        },
    }


async def _task_facets(cur, parser_type: str, where_sql: str, params: list) -> dict:
    facets = _empty_facets()
    bucket_sql = _priority_case(parser_type)
    await cur.execute(
        f"""
        SELECT bucket, COUNT(*)
        FROM (
            SELECT {bucket_sql} AS bucket
            FROM _online_source_projection AS projection
            WHERE {where_sql}
        ) AS priority_rows
        GROUP BY bucket
        """,
        params,
    )
    for bucket, count in await cur.fetchall():
        if str(bucket) in facets["priority_counts"]:
            facets["priority_counts"][str(bucket)] = int(count or 0)
    await cur.execute(
        f"""
        SELECT projection.task_state, COUNT(*)
        FROM _online_source_projection AS projection
        WHERE {where_sql}
        GROUP BY projection.task_state
        """,
        params,
    )
    for state, count in await cur.fetchall():
        if str(state) in facets["status_counts"]:
            facets["status_counts"][str(state)] = int(count or 0)
    if is_registration_task(parser_type):
        await cur.execute(
            f"SELECT COUNT(*) FROM _online_source_projection AS projection "
            f"WHERE {where_sql} AND {_registration_review_condition(parser_type)}",
            params,
        )
        facets["registration_review_count"] = int((await cur.fetchone())[0] or 0)
    if parser_type == MODEL_THREE_PARSER:
        await cur.execute(
            f"""
            SELECT qmf_state,COUNT(*)
            FROM (
                SELECT {_qmf_feedback_state_sql()} AS qmf_state
                FROM _online_source_projection AS projection
                WHERE {where_sql} AND projection.task_state='completed'
            ) AS qmf_rows
            GROUP BY qmf_state
            """,
            params,
        )
        for state, count in await cur.fetchall():
            normalized = str(state or "error")
            if normalized in facets["qmf_feedback_counts"]:
                facets["qmf_feedback_counts"][normalized] = int(count or 0)
    facets["total"] = sum(facets["priority_counts"].values())
    return facets


async def _task_filter_options(
    cur,
    parser_type: str,
    context: dict,
    scope: FlowScope,
    user: dict,
    communities: list[str] | None = None,
    small_communities: list[str] | None = None,
    match_status: list[str] | None = None,
    review_stage: ReviewStage = "all",
) -> dict:
    capability_user = _task_capability_user(user)
    scope_where, scope_params = _scope_where(context, scope)
    where_sql = f"projection.parser_type=%s AND {scope_where}"
    params = [parser_type, *scope_params]
    if review_stage != "all":
        where_sql = f"{where_sql} AND {_review_stage_condition(parser_type, review_stage)}"
    inspector_where = where_sql
    inspector_params = list(params)
    community_condition, community_params = _multi_filter_condition(
        "community", communities or []
    )
    small_community_condition, small_community_params = _multi_filter_condition(
        "small_community_name", small_communities or []
    )
    match_status_condition, match_status_params = _multi_filter_condition(
        "address_match_status", match_status or []
    )
    if community_condition != "1=1":
        inspector_where = f"{inspector_where} AND {community_condition}"
        inspector_params.extend(community_params)
    if small_community_condition != "1=1":
        inspector_where = f"{inspector_where} AND {small_community_condition}"
        inspector_params.extend(small_community_params)
    if match_status_condition != "1=1":
        inspector_where = f"{inspector_where} AND {match_status_condition}"
        inspector_params.extend(match_status_params)
    result = {
        "communities": [], "small_communities": [],
        "match_statuses": [], "inspectors": [], "watch_categories": [],
    }
    for column, key, empty_label in (
        ("community", "communities", "社区未填写"),
        ("small_community_name", "small_communities", "小区未关联"),
        ("address_match_status", "match_statuses", "未匹配"),
        ("inspector", "inspectors", "未分配核查人"),
    ):
        option_where = inspector_where if column == "inspector" else where_sql
        option_params = inspector_params if column == "inspector" else params
        if column == "small_community_name" and communities:
            option_where = f"{option_where} AND {community_condition}"
            option_params = [*option_params, *community_params]
        if column == "address_match_status" and communities:
            option_where = f"{option_where} AND {community_condition}"
            option_params = [*option_params, *community_params]
        await cur.execute(
            f"""
            SELECT projection.{column}, COUNT(*)
            FROM _online_source_projection AS projection
            WHERE {option_where}
            GROUP BY projection.{column}
            ORDER BY CASE WHEN TRIM(COALESCE(projection.{column}, ''))='' THEN 1 ELSE 0 END,
                     projection.{column}
            """,
            option_params,
        )
        for value, count in await cur.fetchall():
            normalized = str(value or "").strip()
            result[key].append({
                "value": normalized or EMPTY_FILTER_VALUE,
                "label": normalized or empty_label,
                "count": int(count or 0),
            })
    if settings.REGISTRY_FEATURE_ENABLED:
        registry = settings.MYSQL_REGISTRY_DB.replace("`", "")
        await cur.execute(
            f"""
            SELECT category.id, category.name, category.color, category.alert_level,
                   COUNT(DISTINCT projection.row_key)
            FROM _online_source_projection projection
            JOIN `{registry}`.online_task_watch_snapshots snapshot
              ON snapshot.parser_type=projection.parser_type
             AND snapshot.row_key=projection.row_key
            JOIN `{registry}`.watch_assignments assignment
              ON assignment.id=snapshot.assignment_id
            JOIN `{registry}`.watch_categories category
              ON category.id=assignment.category_id
            WHERE {where_sql}
            GROUP BY category.id, category.name, category.color, category.alert_level
            ORDER BY category.sort_order, category.id
            """,
            params,
        )
        result["watch_categories"] = [
            {"value": int(row[0]), "label": str(row[1]), "color": str(row[2]),
             "alert_level": str(row[3]), "count": int(row[4] or 0)}
            for row in await cur.fetchall()
        ]
    assignment_enabled = (
        _can_assign_tasks(context)
        and (
            has_permission(user, ONLINE_RAW_EDIT)
            or has_permission(user, ONLINE_TASK_MANAGE)
        )
    )
    assignment_context = (
        await inspector_option_context(cur, capability_user, assignment_only=True)
        if assignment_enabled
        else {
            "community_aliases": {},
            "inspectors_by_community": {},
        }
    )
    result["assignment"] = {
        "enabled": assignment_enabled,
        "community_aliases": assignment_context["community_aliases"],
        "inspectors_by_community": assignment_context["inspectors_by_community"],
    }
    return result


async def _list_mobile_tasks_data(
    parser_type: str,
    data: TaskSearch,
    user: dict,
    conn,
) -> dict:
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    context = await _flow_context(conn, user)
    where_sql, query_params = _task_where(context, parser_type, data)
    facet_data = data.model_copy(update={
        "status": "all",
        "review_stage": "all",
        "priority": "all",
        "qmf_feedback_states": [],
    })
    base_where, base_params = _task_where(
        context,
        parser_type,
        facet_data,
        include_priority=False,
    )
    async with conn.cursor() as cur:
        ready = await _task_source_ready(cur, parser_type)
        if not ready:
            return {
                "data": [],
                "total": 0,
                "page": data.page,
                "page_size": data.page_size,
                "source_ready": False,
                "message": (
                    "本地任务来源尚未建立"
                    if local_data_source_enabled()
                    else "来源定位尚未建立，请等待一次正常同步"
                ),
                "facets": _empty_facets(),
                "priority_labels": PRIORITY_LABELS,
                "filters": {
                    "scope": data.scope,
                    "status": data.status,
                    "review_stage": data.review_stage,
                    "communities": data.communities,
                    "inspectors": data.inspectors,
                    "watch_categories": data.watch_categories,
                    "qmf_feedback_states": data.qmf_feedback_states,
                    "priority": data.priority,
                    "sort": data.sort,
                    "keyword_present": bool(data.keyword.strip()),
                },
            }
        await cur.execute(
            f"SELECT COUNT(*) FROM _online_source_projection AS projection "
            f"WHERE {where_sql}",
            query_params,
        )
        total = int((await cur.fetchone())[0] or 0)
        facets = await _task_facets(cur, parser_type, base_where, base_params)
        order_sql = _task_order(parser_type, data.sort)
        await cur.execute(
            f"""
            SELECT projection.row_key, projection.values_json,
                   projection.source_count, projection.conflict,
                   projection.pending_state, projection.task_state
            FROM _online_source_projection AS projection
            WHERE {where_sql}
            ORDER BY {order_sql}
            LIMIT %s OFFSET %s
            """,
            [*query_params, data.page_size, (data.page - 1) * data.page_size],
        )
        rows = await cur.fetchall()
        address_matches = await _address_matches_by_rows(
            cur, parser_type, [str(row[0]) for row in rows]
        )
        active_counts = await _active_source_counts(
            cur,
            parser_type,
            [str(row[0]) for row in rows],
        )
        qmf_by_row = await _qmf_status_by_rows(cur, parser_type, rows)
        residence_by_row = await residence_status_by_rows(cur, parser_type, rows)
        registration_by_row = await registration_links_by_rows(
            cur,
            parser_type,
            [str(row[0]) for row in rows],
        )
        review_by_row = await review_flows_by_rows(
            cur, [(parser_type, str(row[0])) for row in rows]
        )
        watch_by_row = await task_watch_payload(
            cur,
            parser_type,
            [str(row[0]) for row in rows],
        ) if settings.REGISTRY_FEATURE_ENABLED else {}
    photo_fetched_rows = await _task_photo_fetched_rows(
        user,
        parser_type,
        [str(row[0]) for row in rows],
    )
    return {
        "data": [
            _task_record(
                parser_type,
                str(row[0]),
                json_value(row[1], {}),
                active_counts.get(str(row[0]), 0)
                if local_data_source_enabled()
                else int(row[2] or 0),
                active_counts.get(str(row[0]), 0) > 1
                if local_data_source_enabled()
                else bool(row[3]),
                bool(str(row[4] or "")),
                str(row[5] or ""),
                watch_by_row.get(str(row[0])),
                str(row[0]) in photo_fetched_rows,
                str(row[4] or ""),
                qmf_status=(
                    qmf_by_row.get(str(row[0]))
                    if str(row[5] or "") == "completed"
                    else None
                ),
                residence_status=residence_by_row.get(str(row[0])),
                registration_link=registration_by_row.get(str(row[0])),
                review_flow=review_by_row.get((parser_type, str(row[0]))),
                address_match=address_matches.get(str(row[0])),
            )
            for row in rows
        ],
        "total": total,
        "page": data.page,
        "page_size": data.page_size,
        "source_ready": True,
        "message": "",
        "facets": facets,
        "priority_labels": PRIORITY_LABELS,
        "filters": {
            "scope": data.scope,
            "status": data.status,
            "review_stage": data.review_stage,
            "communities": data.communities,
            "inspectors": data.inspectors,
            "watch_categories": data.watch_categories,
            "qmf_feedback_states": data.qmf_feedback_states,
            "priority": data.priority,
            "sort": data.sort,
            "keyword_present": bool(data.keyword.strip()),
        },
    }


@router.get("/analysis/filter-options")
async def get_mobile_task_analysis_filter_options(
    parser_type: list[str] = Query(default=[]),
    community: list[str] = Query(default=[]),
    review_stage: ReviewStage = Query("all"),
    user: dict = Depends(require_permission(ONLINE_TASK_MANAGE)),
    conn=Depends(get_db),
):
    parser_types = list(dict.fromkeys(parser_type or list(UNVERIFIABLE_REVIEW_TYPES)))
    if any(value not in UNVERIFIABLE_REVIEW_TYPES for value in parser_types):
        raise HTTPException(400, "存在尚未接入研判工作台的业务表")
    context = await _flow_context(conn, user)
    data = AnalysisTaskSearch(
        parser_types=parser_types,
        review_stage=review_stage,
        communities=community,
    )
    async with conn.cursor() as cur:
        ready = all([
            await _task_source_ready(cur, value)
            for value in parser_types
        ])
        if not ready:
            return {
                "source_ready": False,
                "communities": [],
                "inspectors": [],
                "watch_categories": [],
                "assignment": {
                    "enabled": False,
                    "community_aliases": {},
                    "inspectors_by_community": {},
                },
            }
        options = await _analysis_filter_options(cur, context, user, data)
    return {"source_ready": True, **options}


@router.post("/analysis/search")
async def search_mobile_task_analysis(
    data: AnalysisTaskSearch,
    user: dict = Depends(require_permission(ONLINE_TASK_MANAGE)),
    conn=Depends(get_db),
):
    return await _list_analysis_tasks_data(data, user, conn)


@router.post("/analysis/export")
async def export_mobile_task_analysis(
    data: AnalysisTaskSearch,
    request: Request,
    user: dict = Depends(require_permission(ONLINE_TASK_MANAGE)),
    conn=Depends(get_db),
):
    workbook, count = await _mobile_export_workbook(
        data=data, parser_type=None, user=user, conn=conn,
    )
    await record_admin_audit(
        user,
        "mobile_tasks.analysis_export",
        target_type="mobile_task_analysis",
        target_name="研判任务",
        detail={
            "file_format": "XLSX",
            "parser_types": list(dict.fromkeys(data.parser_types)),
            "review_stage": data.review_stage,
            "communities_count": len(data.communities),
            "inspectors_count": len(data.inspectors),
            "sort": data.sort,
            "keyword_present": bool(data.keyword.strip()),
            "rows": count,
        },
        **request_audit_fields(request),
    )
    filename = f"研判任务-{datetime.now():%Y%m%d%H%M%S}.xlsx"
    return StreamingResponse(
        workbook,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/analysis/import")
async def import_mobile_task_analysis(
    request: Request,
    file: UploadFile = File(...),
    user: dict = Depends(require_permission(ONLINE_TASK_MANAGE)),
    conn=Depends(get_db),
):
    """Import a previously exported analysis workbook without creating tasks."""
    content = await file.read()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.values)
    except Exception as exc:
        raise HTTPException(400, "研判文件无法读取，请上传平台导出的 XLSX 文件") from exc
    if not values:
        raise HTTPException(400, "研判文件为空")
    headers = [str(value or "").strip() for value in values[0]]
    required = ["业务类型", "任务标识", "来源ID", "来源版本", "来源行哈希", "流程版本", "研判阶段", "本次研判决定", "研判意见"]
    missing = [header for header in required if header not in headers]
    if missing:
        raise HTTPException(400, f"研判文件缺少字段：{'、'.join(missing)}")
    index = {header: headers.index(header) for header in required}
    succeeded: list[dict] = []
    failed: list[dict] = []
    for row_number, row in enumerate(values[1:], start=2):
        cell = lambda key: str(row[index[key]] or "").strip() if index[key] < len(row) else ""
        outcome = cell("本次研判决定").lower()
        if not outcome:
            continue
        parser_type = cell("业务类型")
        try:
            source_id = int(cell("来源ID"))
            expected_revision = int(cell("来源版本"))
            flow_version = int(cell("流程版本"))
        except ValueError:
            failed.append({"row": row_number, "reason": "来源ID、来源版本和流程版本必须是数字"})
            continue
        normalized_outcome = {"成功": "success", "研判成功": "success", "success": "success", "失败": "failure", "研判失败": "failure", "failure": "failure"}.get(outcome)
        if normalized_outcome is None:
            failed.append({"row": row_number, "reason": "本次研判决定只能填写成功或失败"})
            continue
        stage = {
            INITIAL_PENDING: INITIAL_PENDING,
            UNVERIFIABLE_STATE_LABELS[INITIAL_PENDING]: INITIAL_PENDING,
            DEEP_PENDING: DEEP_PENDING,
            UNVERIFIABLE_STATE_LABELS[DEEP_PENDING]: DEEP_PENDING,
        }.get(cell("研判阶段"))
        if stage is None:
            failed.append({"row": row_number, "reason": "研判阶段不是当前可提交的初步或深度待研判状态"})
            continue
        if not parser_type or not supports_unverifiable_review(parser_type):
            failed.append({"row": row_number, "reason": "该业务不支持通过研判文件提交两级研判"})
            continue
        try:
            decision = UnverifiableDecision(
                stage=stage,
                outcome=normalized_outcome,
                opinion=cell("研判意见"),
                flow_version=flow_version,
                expected_revision=expected_revision,
                expected_row_hash=cell("来源行哈希"),
            )
            result = await decide_mobile_task_unverifiable_review(
                parser_type, source_id, decision, request, user, conn,
            )
            succeeded.append({"row": row_number, "task": f"{parser_type}:{cell('任务标识')}", "state": result.get("review_flow", {}).get("state", "")})
        except (HTTPException, ValueError) as exc:
            reason = exc.detail if isinstance(exc, HTTPException) else str(exc)
            failed.append({"row": row_number, "reason": reason})
    await record_admin_audit(
        user,
        "mobile_tasks.analysis_import",
        target_type="mobile_task_analysis",
        target_name=file.filename or "研判文件",
        detail={"file_format": "XLSX", "success_count": len(succeeded), "failed_count": len(failed)},
        **request_audit_fields(request),
    )
    return {"success_count": len(succeeded), "failed_count": len(failed), "success": succeeded, "failed": failed}


@router.get("/{parser_type}/filter-options")
async def get_mobile_task_filter_options(
    parser_type: str,
    scope: FlowScope = Query("mine"),
    community: list[str] = Query(default=[]),
    small_community: list[str] = Query(default=[]),
    match_status: list[str] = Query(default=[]),
    review_stage: ReviewStage = Query("all"),
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    context = await _flow_context(conn, user)
    async with conn.cursor() as cur:
        if not await _task_source_ready(cur, parser_type):
            return {
                "source_ready": False,
                "communities": [],
                "inspectors": [],
                "watch_categories": [],
                "assignment": {
                    "enabled": False,
                    "community_aliases": {},
                    "inspectors_by_community": {},
                },
            }
        options = await _task_filter_options(
            cur,
            parser_type,
            context,
            scope,
            user,
            communities=community,
            small_communities=small_community,
            match_status=match_status,
            review_stage=review_stage,
        )
    return {"source_ready": True, **options}


@router.post("/{parser_type}/search")
async def search_mobile_tasks(
    parser_type: str,
    data: TaskSearch,
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    return await _list_mobile_tasks_data(parser_type, data, user, conn)


@router.post("/{parser_type}/export")
async def export_mobile_tasks(
    parser_type: str,
    data: TaskSearch,
    request: Request,
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    workbook, count = await _mobile_export_workbook(
        data=data, parser_type=parser_type, user=user, conn=conn,
    )
    await record_admin_audit(
        user,
        "mobile_tasks.export",
        target_type="mobile_task",
        target_name=parser_type,
        detail={
            "file_format": "XLSX",
            "status": data.status,
            "review_stage": data.review_stage,
            "communities_count": len(data.communities),
            "inspectors_count": len(data.inspectors),
            "sort": data.sort,
            "keyword_present": bool(data.keyword.strip()),
            "rows": count,
        },
        **request_audit_fields(request),
    )
    filename = f"{parser_type}-任务-{datetime.now():%Y%m%d%H%M%S}.xlsx"
    return StreamingResponse(
        workbook,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/{parser_type}/assignment-selection")
async def select_mobile_tasks_for_assignment(
    parser_type: str,
    data: TaskSearch,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """Resolve every currently eligible task for the complete filter result."""
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if not _can_assign_tasks(context):
        raise HTTPException(403, "只有组长及有权管理任务的上级岗位可以批量分配核查人")

    where_sql, query_params = _task_where(context, parser_type, data)
    assignment_source_condition = (
        "projection.conflict=0 "
        "AND projection.address_match_status IN ('confirmed','suggested') "
        "AND projection.small_community_id IS NOT NULL "
        "AND TRIM(COALESCE(projection.community,''))<>''"
    )
    if local_data_source_enabled():
        assignment_source_condition = (
            f"{_active_source_count_sql(parser_type)} <= 1 "
            "AND projection.address_match_status IN ('confirmed','suggested') "
            "AND projection.small_community_id IS NOT NULL "
            "AND TRIM(COALESCE(projection.community,''))<>''"
        )
    async with conn.cursor() as cur:
        if not local_data_source_enabled() and not await _writeback_enabled(cur):
            raise HTTPException(503, "在线回写已由超级管理员暂停")
        await cur.execute(
            f"""
            SELECT projection.row_key, projection.community, projection.small_community_name,
                   projection.address_match_status
            FROM _online_source_projection AS projection
            WHERE {where_sql}
              AND TRIM(COALESCE(projection.inspector, ''))=''
              AND projection.task_state<>'completed'
              AND {assignment_source_condition}
              AND EXISTS (
                  SELECT 1 FROM _online_source_rows AS source_row
                  WHERE source_row.parser_type=projection.parser_type
                    AND source_row.row_key=projection.row_key
                    AND source_row.archived_at IS NULL
                    {active_source_sql_filter(parser_type, 'source_row')}
              )
            ORDER BY {_address_order(parser_type)}, projection.row_key
            LIMIT %s
            """,
            [*query_params, MAX_BULK_ASSIGNMENT_TASKS + 1],
        )
        rows = await cur.fetchall()
        if len(rows) > MAX_BULK_ASSIGNMENT_TASKS:
            raise HTTPException(
                409,
                f"当前筛选可分配任务超过 {MAX_BULK_ASSIGNMENT_TASKS} 条，请继续缩小筛选范围",
            )
        assignment_context = await inspector_option_context(
            cur,
            user,
            assignment_only=True,
        )

    aliases = assignment_context["community_aliases"]
    formal_communities = {
        aliases.get(str(community or "").strip(), "")
        for _, community, _, _ in rows
    }
    if "" in formal_communities:
        raise HTTPException(403, "部分任务社区不在当前账号可分配范围内")
    if len(formal_communities) > 1:
        raise HTTPException(400, "请先筛选到一个社区，再全选当前筛选结果")
    formal_community = next(iter(formal_communities), "")
    if formal_community and not assignment_context["inspectors_by_community"].get(
        formal_community
    ):
        raise HTTPException(400, "该社区当前没有在岗组员可分配")
    return {
        "row_keys": [str(row[0]) for row in rows],
        "total": len(rows),
        "community": formal_community,
    }


@router.post("/{parser_type}/{row_key}/address-match/confirm")
async def confirm_mobile_task_address_match(
    parser_type: str,
    row_key: str,
    data: AddressMatchConfirm,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """人工确认任务的小区归属；确认结果不会被后续规则重跑覆盖。"""
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if not _can_assign_tasks(context):
        raise HTTPException(403, "只有组长及有权管理任务的上级岗位可以确认小区归属")
    scope_where, scope_params = _scope_where(
        context,
        "all" if context.get("admin_mode") else "community",
    )
    await conn.begin()
    try:
        async with conn.cursor() as cur:
            await cur.execute(
                f"""
                SELECT projection.community, projection.source_count,
                       projection.conflict, address_match.original_address
                FROM _online_source_projection AS projection
                LEFT JOIN _online_task_address_matches AS address_match
                  ON address_match.parser_type=projection.parser_type
                 AND address_match.row_key=projection.row_key
                WHERE projection.parser_type=%s AND projection.row_key=%s
                  AND {scope_where}
                FOR UPDATE
                """,
                (parser_type, row_key, *scope_params),
            )
            projection = await cur.fetchone()
            if not projection:
                raise HTTPException(404, "任务不存在或不在当前账号可操作范围内")
            if bool(projection[2]) or int(projection[1] or 0) != 1:
                raise HTTPException(409, "任务存在重复或冲突来源，暂不能确认小区")
            await cur.execute(
                """
                SELECT entry.id, entry.name, entry.community_id, community.name
                FROM _police_address_entries AS entry
                JOIN _communities AS community ON community.id=entry.community_id
                WHERE entry.id=%s AND entry.enabled=1
                """,
                (data.small_community_id,),
            )
            entry = await cur.fetchone()
            if not entry:
                raise HTTPException(404, "小区不存在、已停用或尚未设置所属社区")
            assignment_context = await inspector_option_context(
                cur,
                user,
                assignment_only=True,
            )
            aliases = assignment_context["community_aliases"]
            task_community = aliases.get(str(projection[0] or "").strip(), "")
            entry_community = aliases.get(str(entry[3] or "").strip(), "")
            if not task_community or not entry_community or task_community != entry_community:
                raise HTTPException(409, "所选小区与任务正式社区不一致，请先处理地址冲突")
            await record_feedback_confirmation(
                cur,
                parser_type=parser_type,
                row_key=row_key,
                address=str(projection[3] or ""),
                community_name=str(entry[3] or ""),
                community_id=int(entry[2]),
                confirmed_entry_id=int(entry[0]),
                confirmed_by=int(user["id"]),
            )
            await cur.execute(
                """
                INSERT INTO _online_task_address_matches (
                    parser_type, row_key, original_address,
                    suggested_entry_id, suggested_community_id,
                    suggested_community_name, match_status, match_score,
                    match_method, match_reason, candidates_json,
                    matcher_version, confirmed_entry_id, confirmed_by,
                    confirmed_at
                )
                SELECT parser_type, row_key, '', %s, %s, %s,
                       'confirmed', 1, '人工确认', '管理员已确认小区归属',
                       JSON_ARRAY(), %s, %s, %s, UTC_TIMESTAMP()
                FROM _online_source_projection
                WHERE parser_type=%s AND row_key=%s
                ON DUPLICATE KEY UPDATE
                    suggested_entry_id=VALUES(suggested_entry_id),
                    suggested_community_id=VALUES(suggested_community_id),
                    suggested_community_name=VALUES(suggested_community_name),
                    match_status='confirmed', match_score=1,
                    match_method='人工确认',
                    match_reason='管理员已确认小区归属',
                    matcher_version=VALUES(matcher_version),
                    confirmed_entry_id=VALUES(confirmed_entry_id),
                    confirmed_by=VALUES(confirmed_by),
                    confirmed_at=VALUES(confirmed_at)
                """,
                (
                    int(entry[0]), int(entry[2]), str(entry[3] or ""),
                    MATCHER_VERSION, int(entry[0]), int(user["id"]),
                    parser_type, row_key,
                ),
            )
            await rebuild_projection_rows(cur, parser_type, [row_key])
            result = (await _address_matches_by_rows(cur, parser_type, [row_key])).get(row_key)
        await conn.commit()
    except Exception:
        await conn.rollback()
        raise
    await record_admin_audit(
        user,
        "mobile_tasks.address_match_confirm",
        target_type="mobile_task_address_match",
        target_name=f"{parser_type}:{row_key}",
        detail={"small_community_id": int(entry[0])},
        **request_audit_fields(request),
    )
    return {"message": "小区归属已确认", "address_match": result}


@router.post("/{parser_type}/{row_key}/address-match/resolve-conflict")
async def resolve_mobile_task_address_conflict(
    parser_type: str,
    row_key: str,
    data: AddressMatchConflictResolution,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """提供地址冲突的可执行处理路径，并在修正后重新生成匹配结果。"""
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    if not local_data_source_enabled():
        raise HTTPException(409, "本地业务数据源尚未启用，暂不能处理地址冲突")
    scoped_user = _require_unverifiable_reviewer(user)
    context = await _flow_context(conn, scoped_user)
    scope_where, scope_params = _scope_where(
        context,
        "all" if context.get("admin_mode") else "community",
    )
    parser = get_parser(parser_type)
    community_field = parser.COMMUNITY_COLUMN
    async with conn.cursor() as cur:
        await cur.execute(
            f"""
            SELECT source.id, source.row_key, source.revision, source.row_hash,
                   source.values_json, projection.address_match_status,
                   projection.address_match_candidates, projection.community,
                   projection.source_count, projection.conflict
            FROM _online_source_rows AS source
            JOIN _online_source_projection AS projection
              ON projection.parser_type=source.parser_type
             AND projection.row_key=source.row_key
            WHERE source.id=(
                SELECT candidate_source.id FROM _online_source_rows AS candidate_source
                WHERE candidate_source.parser_type=%s
                  AND candidate_source.row_key=%s
                  AND candidate_source.archived_at IS NULL
                {active_source_sql_filter(parser_type, 'candidate_source')}
                ORDER BY id LIMIT 1
            )
              AND source.parser_type=%s
              AND source.archived_at IS NULL
              AND {scope_where}
            """,
            (parser_type, row_key, parser_type, *scope_params),
        )
        source = await cur.fetchone()
        if not source:
            raise HTTPException(404, "任务不存在或不在当前账号可处理范围内")
        if str(source[1] or "") != row_key:
            raise HTTPException(409, "任务主键已变化，请刷新后重试")
        if int(source[2] or 0) != data.expected_revision or str(source[3] or "") != data.expected_row_hash:
            raise HTTPException(409, "任务已被其他人更新，请刷新后重试")
        if int(source[8] or 0) != 1 or bool(source[9]):
            raise HTTPException(409, "任务存在重复或冲突来源，请先处理来源异常")
        if str(source[5] or "") != "conflict":
            raise HTTPException(409, "该任务当前已不处于地址冲突状态，请刷新后重试")
        candidates = json_value(source[6], [])
        selected = next(
            (item for item in candidates if int(item.get("entry_id") or 0) == data.small_community_id),
            None,
        )
        if not selected:
            raise HTTPException(400, "所选小区不是本次冲突匹配生成的候选，请刷新后重新选择")
        await cur.execute(
            """
            SELECT entry.id, entry.name, entry.community_id, community.name
            FROM _police_address_entries AS entry
            JOIN _communities AS community ON community.id=entry.community_id
            WHERE entry.id=%s AND entry.enabled=1 AND community.is_active=1
            """,
            (data.small_community_id,),
        )
        entry = await cur.fetchone()
        if not entry:
            raise HTTPException(404, "候选小区不存在、已停用或所属社区已停用")
        target_community = str(entry[3] or "").strip()
        if not target_community:
            raise HTTPException(409, "候选小区尚未设置有效所属社区，请先维护小区地址库")
        current_values = json_value(source[4], {})

    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=int(source[0]),
        changes={community_field: target_community},
        base_values={community_field: str(current_values.get(community_field) or "")},
        expected_revision=data.expected_revision,
        request=request,
        user=scoped_user,
        conn=conn,
        explicit_text_edit=True,
        allowed_columns={community_field},
        audit_action="address_match.resolve_conflict",
        record_unverifiable_save=False,
    )
    await record_admin_audit(
        scoped_user,
        "mobile_tasks.address_match.resolve_conflict",
        target_type="mobile_task_address_match",
        target_name=f"{parser_type}:{row_key}",
        detail={
            "small_community_id": int(entry[0]),
            "target_community": target_community,
        },
        **request_audit_fields(request),
    )
    return {
        **result,
        "message": f"已将任务社区修正为{target_community}，正在重新匹配",
        **await _mobile_task_save_context(conn, parser_type, str(result["row_key"])),
    }


@router.get("/{parser_type}/assignment-workbench")
async def get_mobile_task_assignment_workbench(
    parser_type: str,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """Return an indexed, address-sorted queue of unassigned tasks.

    The legacy expression ``ORDER BY {_address_order(parser_type)}`` is kept
    out of the SQL path; the projection now stores its normalized sort key.
    """
    started_at = time.monotonic()
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if not _can_assign_tasks(context):
        raise HTTPException(403, "只有组长及有权管理任务的上级岗位可以批量分配核查人")
    scope_where, scope_params = _scope_where(
        context,
        "all" if context.get("admin_mode") else "community",
    )
    assignment_source_condition = "projection.assignment_queue_ready=1"
    async with conn.cursor() as cur:
        if not local_data_source_enabled() and not await _writeback_enabled(cur):
            raise HTTPException(503, "在线回写已由超级管理员暂停")
        await cur.execute(
            f"""
            SELECT COUNT(*)
            FROM _online_source_projection AS projection
            WHERE projection.parser_type=%s
              AND {scope_where}
              AND {assignment_source_condition}
            """,
            [parser_type, *scope_params],
        )
        available_total = int((await cur.fetchone())[0] or 0)
        await cur.execute(
            f"""
            SELECT projection.row_key, projection.community,
                   projection.small_community_name, projection.address_match_status,
                   projection.assignment_source_label,
                   projection.assignment_address_display
            FROM _online_source_projection AS projection
            WHERE projection.parser_type=%s
              AND {scope_where}
              AND {assignment_source_condition}
            ORDER BY projection.assignment_address_sort_key, projection.row_key
            LIMIT %s
            """,
            [parser_type, *scope_params, MAX_BULK_ASSIGNMENT_TASKS],
        )
        rows = await cur.fetchall()
        assignment_context = await inspector_option_context(
            cur,
            user,
            assignment_only=True,
        )
        await cur.execute(
            f"""
            SELECT projection.community, projection.inspector, COUNT(*)
            FROM _online_source_projection AS projection
            WHERE projection.parser_type=%s
              AND {scope_where}
              AND TRIM(COALESCE(projection.inspector, ''))<>''
              AND projection.task_state<>'completed'
              AND projection.source_count=1
              AND projection.conflict=0
            GROUP BY projection.community, projection.inspector
            """,
            [parser_type, *scope_params],
        )
        assigned_count_rows = await cur.fetchall()

    aliases = assignment_context["community_aliases"]
    inspectors_by_community = assignment_context["inspectors_by_community"]
    items: list[dict[str, str]] = []
    for row_key, raw_community, raw_small_community, raw_match_status, raw_source, raw_address in rows:
        community = aliases.get(str(raw_community or "").strip(), "")
        if not community:
            community = str(raw_community or "").strip() or "社区未识别"
        items.append({
            "row_key": str(row_key),
            "community": community,
            "small_community": str(raw_small_community or ""),
            "match_status": str(raw_match_status or "unmatched"),
            "source": str(raw_source or TASK_WORKFLOWS[parser_type].label),
            "address": str(raw_address or "未填写地址"),
        })
    community_counts: dict[str, int] = {}
    for item in items:
        community_counts[item["community"]] = community_counts.get(item["community"], 0) + 1
    inspector_counts_by_community: dict[str, dict[str, int]] = {}
    for raw_community, raw_inspector, count in assigned_count_rows:
        formal_community = aliases.get(str(raw_community or "").strip(), "")
        if not formal_community:
            formal_community = str(raw_community or "").strip() or "社区未识别"
        inspector_name = str(raw_inspector or "").strip()
        if not inspector_name:
            continue
        inspector_counts_by_community.setdefault(formal_community, {})[inspector_name] = int(count or 0)
    assigned_totals_by_community: dict[str, int] = {}
    for formal_community, counts in inspector_counts_by_community.items():
        assigned_totals_by_community[formal_community] = sum(counts.values())
    all_communities = set(community_counts) | set(assigned_totals_by_community)
    return {
        "data": items,
        "total": available_total,
        "displayed_total": len(items),
        "limited": available_total > len(items),
        "limit": MAX_BULK_ASSIGNMENT_TASKS,
        "communities": [
            {"value": community, "label": community, "count": community_counts.get(community, 0),
             "assigned_count": assigned_totals_by_community.get(community, 0)}
            for community in sorted(all_communities)
        ],
        "inspectors_by_community": {
            community: list(inspectors_by_community.get(community) or [])
            for community in all_communities
        },
        "inspector_counts_by_community": {
            community: {
                inspector: inspector_counts_by_community.get(community, {}).get(inspector, 0)
                for inspector in inspectors_by_community.get(community) or []
            }
            for community in all_communities
        },
        "assigned_totals_by_community": assigned_totals_by_community,
        "duration_ms": max(0, int((time.monotonic() - started_at) * 1000)),
        "query_mode": "indexed_projection",
    }


@router.get("/{parser_type}")
async def list_mobile_tasks(
    parser_type: str,
    scope: FlowScope = Query("mine"),
    status: TaskStatus = Query("pending"),
    review_stage: ReviewStage = Query("all"),
    community: list[str] = Query(default=[]),
    small_community: list[str] = Query(default=[]),
    match_status: list[str] = Query(default=[]),
    inspector: list[str] = Query(default=[]),
    watch_category: list[int] = Query(default=[]),
    priority: Priority = Query("all"),
    sort: SortMode = Query("priority"),
    keyword: str | None = Query(None, max_length=100),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=50),
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    return await _list_mobile_tasks_data(
        parser_type,
        TaskSearch(
            scope=scope,
            status=status,
            review_stage=review_stage,
            communities=community,
            small_communities=small_community,
            match_status=match_status,
            inspectors=inspector,
            watch_categories=watch_category,
            priority=priority,
            sort=sort,
            keyword=keyword or "",
            page=page,
            page_size=page_size,
        ),
        user,
        conn,
    )


async def _mobile_task_detail_data(
    parser_type: str,
    row_key: str,
    user: dict,
    conn,
    *,
    analysis_mode: bool = False,
    include_photo_requests: bool = True,
) -> dict:
    if analysis_mode:
        user = (
            _require_unverifiable_reviewer(user)
            if supports_unverifiable_review(parser_type)
            else _require_analysis_user(user)
        )
    capability_user = _task_capability_user(user)
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    context = await _flow_context(conn, user)
    parser = get_parser(parser_type)
    dependency_blocked = False
    detail_scope: FlowScope = "all" if context["admin_mode"] else "community"
    scope_where, scope_params = _scope_where(context, detail_scope)
    async with conn.cursor() as cur:
        if not await _task_source_ready(cur, parser_type):
            raise HTTPException(
                409,
                "本地任务来源尚未建立"
                if local_data_source_enabled()
                else "来源定位尚未建立，请等待一次正常同步",
            )
        await cur.execute(
            f"""
            SELECT values_json, source_count, conflict, pending_state, task_state
            FROM _online_source_projection AS projection
            WHERE projection.parser_type=%s AND projection.row_key=%s
              AND {scope_where}
            """,
            (parser_type, row_key, *scope_params),
        )
        parent_row = await cur.fetchone()
        if not parent_row:
            raise HTTPException(404, "任务不存在或不属于当前社区")
        parent_values = json_value(parent_row[0], {})
        review_by_row = await review_flows_by_rows(cur, [(parser_type, row_key)])
        review_flow = review_by_row.get((parser_type, row_key))
        if review_flow:
            review_flow["events"] = await review_events_for_flow(cur, int(review_flow["id"]))
        if not analysis_mode:
            dependency_blocked = await online_task_blocked(cur, parser_type, row_key)
        if analysis_mode and not (
            review_flow and review_flow.get("state") in {
                INITIAL_PENDING, INITIAL_EXTENSION, DEEP_PENDING, DEEP_EXTENSION,
            }
        ) and TASK_WORKFLOWS[parser_type].review_stage(parent_values) not in {
            "waiting_analysis", "analyzed"
        }:
            raise HTTPException(404, "该任务当前不属于网格核查研判范围")
        await cur.execute(
            f"""
            SELECT id, physical_row, values_json, cell_meta_json,
                   revision, row_hash, spreadsheet_id, sheet_id
            FROM _online_source_rows AS source
            WHERE source.parser_type=%s AND source.row_key=%s
              AND source.archived_at IS NULL
              {active_source_sql_filter(parser_type)}
            ORDER BY spreadsheet_id, physical_row
            """,
            (parser_type, row_key),
        )
        raw_sources = await cur.fetchall()
        local_changes = await load_local_changes(
            cur, [int(row[0]) for row in raw_sources]
        )
        enabled = local_data_source_enabled() or await _writeback_enabled(cur)
        assignment_context = (
            await inspector_option_context(cur, capability_user, assignment_only=True)
            if _can_assign_tasks(context) and (
                has_permission(user, ONLINE_RAW_EDIT)
                or has_permission(user, ONLINE_TASK_MANAGE)
            )
            else None
        )
        watch_by_row = await task_watch_payload(cur, parser_type, [row_key]) \
            if settings.REGISTRY_FEATURE_ENABLED else {}
        qmf_by_row = await _qmf_status_by_rows(
            cur,
            parser_type,
            [(row_key, parent_row[0])],
        )
        residence_by_row = await residence_status_by_rows(
            cur,
            parser_type,
            [(row_key, parent_row[0])],
        )
        registration_by_row = await registration_links_by_rows(
            cur,
            parser_type,
            [row_key],
        )
        registration_link = registration_by_row.get(row_key)
        address_match = (await _address_matches_by_rows(
            cur,
            parser_type,
            [row_key],
        )).get(row_key)

        sources = []
        for (
            source_id,
            physical_row,
            raw_values,
            raw_meta,
            revision,
            row_hash,
            spreadsheet_id,
            sheet_id,
        ) in raw_sources:
            source_changes = local_changes.get(int(source_id), [])
            values = overlay_local_values(
                json_value(raw_values, {}), source_changes
            )
            # 同一业务主键偶尔会跨社区重复。父投影可以用于定位任务，
            # 但详情绝不能因此暴露其他社区的腾讯原始行。
            if not _source_in_community(
                parser, values, context["community_values"]
            ):
                continue
            capabilities = await row_edit_capabilities(
                cur, capability_user, parser, values
            )
            metadata = await _managed_column_metadata(
                cur,
                parser,
                json_value(raw_meta, {}),
                spreadsheet_id=int(spreadsheet_id),
                sheet_id=str(sheet_id),
            )
            workflow = TASK_WORKFLOWS[parser_type]
            if is_registration_task(parser_type) and workflow.result_field in parser.COLUMNS:
                metadata[workflow.result_field] = {
                    "type": "select",
                    "multiple": False,
                    "options": [
                        {"id": option, "text": option}
                        for option in workflow.result_options
                        if option != "已登记"
                    ],
                }
            raw_community = parser.community_value(values)
            formal_community = (
                (assignment_context or {}).get("community_aliases", {})
                .get(str(raw_community or "").strip(), "")
            )
            assignee_options = list(
                (assignment_context or {}).get("inspectors_by_community", {})
                .get(formal_community, [])
            )
            if "核查人" in metadata and assignee_options:
                metadata["核查人"] = {
                    "type": "select",
                    "multiple": False,
                    "options": [
                        {"id": name, "text": name}
                        for name in assignee_options
                    ],
                }
            source_task = _task_record(
                parser_type,
                row_key,
                values,
                1,
                False,
                bool(str(parent_row[3] or "")),
                TASK_WORKFLOWS[parser_type].state(
                    values,
                    registration_status=str(
                        (registration_link or {}).get("status") or ""
                    ),
                ),
                review_flow=review_flow,
                address_match=address_match,
            )
            editable_fields = (
                list(capabilities["editable_fields"])
                if enabled else []
            )
            if context["position"] != "组长" and not context["admin_mode"]:
                editable_fields = [
                    field for field in editable_fields if field != "核查人"
                ]
            # 二次反馈是否真正允许保存，仍由批量校验根据提交后的主结果
            # 决定；这里先把字段交给手机表单，才能在同一次提交中展开。
            if enabled and capabilities["can_edit"]:
                editable_fields.extend(
                    field for field in TASK_WORKFLOWS[parser_type].secondary_fields
                    if field in parser.COLUMNS
                )
            analysis_fields = set(TASK_WORKFLOWS[parser_type].analysis_fields)
            if analysis_mode:
                if supports_unverifiable_review(parser_type):
                    # 两级研判只能通过带流程版本、来源版本和行哈希的
                    # 结构化决定接口提交，禁止旧客户端继续自由编辑“研判”列。
                    editable_fields = []
                else:
                    editable_fields = [
                        field for field in analysis_fields
                        if field in parser.COLUMNS
                    ] if source_task["review_stage"] in {
                        "waiting_analysis", "analyzed",
                    } else []
            else:
                editable_fields = [
                    field for field in editable_fields
                    if field not in analysis_fields
                ]
            sources.append({
                "id": int(source_id),
                "row_key": row_key,
                "physical_row": int(physical_row),
                "source_available": int(physical_row) > 0,
                "values": {
                    column: str(values.get(column, "") or "")
                    for column in parser.COLUMNS
                },
                "cell_meta": metadata,
                "revision": int(revision),
                "row_hash": str(row_hash),
                "editable_fields": list(dict.fromkeys(editable_fields)),
                "state": source_task["state"],
                "needs_review": source_task["needs_review"],
                "review_stage": source_task["review_stage"],
                "review_flow": review_flow,
                "sync_state": local_sync_state(source_changes),
                "sync_fields": [
                    {
                        "field": item["field_name"],
                        "platform_value": item["local_value"],
                        "tencent_value": item["remote_value"],
                        "status": item["status"],
                        "error_code": item["error_code"],
                    }
                    for item in source_changes
                ],
            })

        if not sources:
            raise HTTPException(404, "任务不存在或不属于当前社区")

    workflow = TASK_WORKFLOWS[parser_type]
    photo_requests = (
        await _task_photo_results(user, parser_type, row_key)
        if include_photo_requests else []
    )
    qmf_config = await load_qmf_config(conn)
    qmf_allowed = has_permission(user, QMF_REGISTRATION_EXECUTE)
    effective_source_count = len(sources)
    effective_conflict = (
        effective_source_count > 1
        if local_data_source_enabled()
        else bool(parent_row[2])
    )
    qmf_preview = preview_capability(
        allowed=qmf_allowed,
        parser_type=parser_type,
        source_count=effective_source_count,
        conflict=effective_conflict,
        values=sources[0]["values"] if effective_source_count == 1 else None,
        config=qmf_config,
    )
    qmf_registration = registration_capability(
        allowed=qmf_allowed,
        parser_type=parser_type,
        source_count=effective_source_count,
        conflict=effective_conflict,
        values=sources[0]["values"] if effective_source_count == 1 else None,
        config=qmf_config,
    )
    latest_qmf_run, qmf_feedback = await _qmf_registration_state(
        conn,
        parser_type=parser_type,
        sources=sources,
        user=user,
    )
    qmf_registration["latest_run"] = latest_qmf_run
    return {
        "task": _task_record(
            parser_type,
            row_key,
            parent_values,
            effective_source_count,
            effective_conflict,
            bool(str(parent_row[3] or "")),
            str(parent_row[4] or ""),
            watch_by_row.get(row_key),
            False,
            str(parent_row[3] or ""),
            qmf_status=(
                qmf_by_row.get(row_key)
                if str(parent_row[4] or "") == "completed"
                else None
            ),
            residence_status=residence_by_row.get(row_key),
            registration_link=registration_link,
            review_flow=review_flow,
            address_match=address_match,
        ),
        "workflow": {
            "label": workflow.label,
            "result_field": workflow.result_field,
            # Keep the canonical field first while allowing historical source
            # aliases in the detail page to match the list summary.
            "phone_fields": list(dict.fromkeys((*workflow.phone_fields, *PHONE_FIELD_ALIASES))),
            "title_fields": list(workflow.title_fields),
            "address_fields": list(workflow.address_fields),
            "date_fields": list(workflow.date_fields),
            "identity_fields": list(workflow.identity_fields),
            "source_fields": list(workflow.source_fields),
            "secondary_fields": list(workflow.secondary_fields),
            "extra_edit_fields": list(getattr(parser, "MOBILE_EDITABLE_FIELDS", ())),
            "analysis_fields": list(workflow.analysis_fields),
            "columns": parser.COLUMNS,
        },
        "writeback_enabled": enabled,
        "analysis_mode": analysis_mode,
        "dependency_blocked": dependency_blocked,
        "dependency_message": (
            "该任务已进入基础管控研判队列，网格员仍可继续核查并修改结果"
            if dependency_blocked else ""
        ),
        "photo_requests": photo_requests,
        "qmf_preview": qmf_preview,
        "qmf_registration": qmf_registration,
        "qmf_feedback": qmf_feedback,
        "qmf_status": qmf_by_row.get(row_key),
        "residence_status": residence_by_row.get(row_key),
        "registration_link": registration_link,
        "address_match": address_match,
        "review_flow": review_flow,
        "registration_manual_confirm_allowed": can_confirm_registration(user),
        "sources": sources,
        "data_source_mode": "local" if local_data_source_enabled() else "tencent",
    }


async def _mobile_task_inline_editors_data(
    parser_type: str,
    row_keys: list[str],
    user: dict,
    conn,
    *,
    analysis_mode: bool = False,
) -> dict:
    """Prepare the current page's editors in one browser request.

    The detail permission, community scope and editable-field calculation stay
    identical to the normal task detail endpoint.  A row that moved out of the
    current scope is returned as unavailable instead of failing the whole page.
    """
    normalized_keys = list(dict.fromkeys(
        str(row_key).strip() for row_key in row_keys if str(row_key).strip()
    ))
    if not normalized_keys:
        raise HTTPException(400, "请至少选择一条任务")

    items: dict[str, dict] = {}
    for row_key in normalized_keys:
        try:
            detail = await _mobile_task_detail_data(
                parser_type,
                row_key,
                user,
                conn,
                analysis_mode=analysis_mode,
                include_photo_requests=False,
            )
        except HTTPException as exc:
            if exc.status_code not in {404, 409}:
                raise
            items[row_key] = {
                "available": False,
                "reason": str(exc.detail),
            }
            continue
        if len(detail["sources"]) != 1 or (
            not local_data_source_enabled() and detail["task"]["conflict"]
        ):
            items[row_key] = {
                "available": False,
                "reason": (
                    "该任务包含多个本地来源，请进入详情处理"
                    if local_data_source_enabled()
                    else "该任务包含多个腾讯来源，请进入详情选择来源后修改"
                ),
            }
            continue
        items[row_key] = {
            "available": True,
            "detail": detail,
        }
    return {"items": items, "analysis_mode": analysis_mode}


@router.post("/analysis/{parser_type}/inline-editors")
async def get_mobile_task_analysis_inline_editors(
    parser_type: str,
    data: InlineEditorRequest,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    return await _mobile_task_inline_editors_data(
        parser_type,
        data.row_keys,
        user,
        conn,
        analysis_mode=True,
    )


@router.post("/{parser_type}/inline-editors")
async def get_mobile_task_inline_editors(
    parser_type: str,
    data: InlineEditorRequest,
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    return await _mobile_task_inline_editors_data(
        parser_type,
        data.row_keys,
        user,
        conn,
    )


@router.get("/analysis/{parser_type}/{row_key}")
async def get_mobile_task_analysis_detail(
    parser_type: str,
    row_key: str,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    return await _mobile_task_detail_data(
        parser_type, row_key, user, conn, analysis_mode=True
    )


@router.patch("/analysis/{parser_type}/source-rows/{source_id}/decision")
async def decide_mobile_task_unverifiable_review(
    parser_type: str,
    source_id: int,
    data: UnverifiableDecision,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """提交结构化初步/深度研判决定，并原子写入延时截止日期。"""
    if not supports_unverifiable_review(parser_type):
        raise HTTPException(400, "该业务不支持无法核实两级研判")
    scoped_user = _require_unverifiable_reviewer(user)
    workflow = TASK_WORKFLOWS[parser_type]
    analysis_field = next(
        (field for field in workflow.analysis_fields if field), ""
    )
    secondary_field = next(
        (field for field in workflow.secondary_fields if field), ""
    )
    deadline_field = workflow.date_fields[0] if workflow.date_fields else ""
    if not analysis_field:
        raise HTTPException(400, "该业务没有可填写的研判字段")
    prepared_holder: dict = {}

    def validate_source(current: dict) -> None:
        if not is_unverifiable_result(current.get(workflow.result_field)):
            raise HTTPException(409, "该任务当前已不是无法核实，请刷新后重试")

    async def transaction_prepare(*, cur, source, current_values, changes):
        del changes
        if int(source.get("revision") or 0) != data.expected_revision:
            raise HTTPException(409, "腾讯来源版本已经变化，请刷新后重试")
        prepared = await prepare_decision(
            cur,
            parser_type=parser_type,
            source=source,
            current_values=current_values,
            stage=data.stage,
            outcome=data.outcome,
            opinion=data.opinion,
            expected_flow_version=data.flow_version,
            expected_row_hash=data.expected_row_hash,
        )
        prepared_holder["value"] = prepared
        result = {analysis_field: prepared["summary"]}
        if prepared["due_date"] and deadline_field:
            result[deadline_field] = prepared["due_date"].isoformat()
        if prepared["next_state"] in {
            INITIAL_EXTENSION, DEEP_PENDING, DEEP_EXTENSION,
        } and secondary_field:
            # 腾讯字段只展示当前阶段的反馈；上一阶段原文已由结构化事件保留。
            result[secondary_field] = ""
        return result

    async def transaction_callback(*, cur, source, before, after, row_key_before, row_key_after, revision):
        del before, after, row_key_before, row_key_after
        prepared = prepared_holder.get("value")
        if not prepared:
            raise HTTPException(409, "研判决定未准备完成，请刷新后重试")
        prepared_holder["applied"] = await apply_decision(
            cur,
            prepared=prepared,
            source=source,
            revision=revision,
            actor_user_id=int(scoped_user.get("id") or 0),
        )

    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=source_id,
        changes={analysis_field: data.opinion.strip()},
        base_values={analysis_field: ""},
        expected_revision=data.expected_revision,
        request=request,
        user=scoped_user,
        conn=conn,
        explicit_text_edit=True,
        allowed_columns={
            field for field in (analysis_field, deadline_field, secondary_field)
            if field
        },
        current_values_validator=validate_source,
        redact_audit_values=True,
        audit_action="unverifiable_review_decision",
        transaction_prepare=transaction_prepare,
        transaction_callback=transaction_callback,
        record_unverifiable_save=False,
    )
    return {**result, "review_flow": prepared_holder.get("applied")}


@router.patch("/analysis/{parser_type}/source-rows/{source_id}")
async def update_mobile_task_analysis(
    parser_type: str,
    source_id: int,
    data: TaskBatchUpdate,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    if supports_unverifiable_review(parser_type):
        raise HTTPException(410, "无法核实任务请使用结构化研判决定接口")
    scoped_user = _require_analysis_user(user)
    parser = get_parser(parser_type)
    analysis_fields = {
        field for field in TASK_WORKFLOWS[parser_type].analysis_fields
        if field in parser.COLUMNS
    }
    if not analysis_fields:
        raise HTTPException(400, "该业务没有可填写的研判字段")
    if not data.changes or any(field not in analysis_fields for field in data.changes):
        raise HTTPException(400, "研判入口只能修改研判字段")
    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT values_json FROM _online_source_rows "
            "WHERE id=%s AND parser_type=%s",
            (source_id, parser_type),
        )
        row = await cur.fetchone()
        if not row:
            raise HTTPException(404, "腾讯来源行不存在")
        values = json_value(row[0], {})
        if TASK_WORKFLOWS[parser_type].review_stage(values) not in {
            "waiting_analysis", "analyzed"
        }:
            raise HTTPException(409, "该任务当前不属于网格核查研判范围")
    def validate_analysis_source(current: dict) -> None:
        if TASK_WORKFLOWS[parser_type].review_stage(current) not in {
            "waiting_analysis", "analyzed"
        }:
            raise HTTPException(409, "该任务当前不属于网格核查研判范围")

    return await queue_source_fields(
        parser_type=parser_type,
        source_id=source_id,
        changes=data.changes,
        base_values=data.base_values,
        expected_revision=data.expected_revision,
        request=request,
        user=scoped_user,
        conn=conn,
        explicit_text_edit=True,
        allowed_columns=analysis_fields,
        current_values_validator=validate_analysis_source,
        redact_audit_values=True,
    )


@router.get("/{parser_type}/{row_key}/residence-detail")
async def get_mobile_task_residence_detail(
    parser_type: str,
    row_key: str,
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    detail = await _mobile_task_detail_data(
        parser_type,
        row_key,
        user,
        conn,
        include_photo_requests=False,
    )
    residence_status = detail.get("residence_status") or {}
    if residence_status.get("state") != "registered":
        raise HTTPException(409, "该任务尚无可展示的居住证登记资料")
    sources = detail.get("sources") or []
    if detail["task"].get("conflict") or len(sources) != 1:
        raise HTTPException(409, "该任务来源不唯一，暂不能读取居住证人员资料")
    source = sources[0]
    if not source.get("source_available"):
        raise HTTPException(409, "该任务来源已变化，请刷新后重试")
    try:
        return await residence_detail_for_values(
            conn,
            parser_type,
            source.get("values") or {},
        )
    except ResidencePlatformError as exc:
        status_code = 409 if exc.code in {
            "detail_not_registered",
            "invalid_identity",
            "community_missing",
            "community_not_found",
            "community_ambiguous",
            "community_code_missing",
            "session_not_ready",
        } else 502
        raise HTTPException(status_code, str(exc)) from exc


@router.get("/{parser_type}/{row_key}")
async def get_mobile_task_detail(
    parser_type: str,
    row_key: str,
    user: dict = Depends(require_permission(ONLINE_RAW_VIEW)),
    conn=Depends(get_db),
):
    return await _mobile_task_detail_data(parser_type, row_key, user, conn)


async def _mobile_task_save_context(conn, parser_type: str, row_key: str) -> dict:
    """Load only the task state changed by an incremental save."""
    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT values_json,source_count,conflict,pending_state,task_state "
            "FROM _online_source_projection WHERE parser_type=%s AND row_key=%s",
            (parser_type, row_key),
        )
        projection = await cur.fetchone()
        if not projection:
            return {}
        values = json_value(projection[0], {})
        registration_link = (await registration_links_by_rows(
            cur,
            parser_type,
            [row_key],
        )).get(row_key)
        review_flow = (await review_flows_by_rows(
            cur,
            [(parser_type, row_key)],
        )).get((parser_type, row_key))
        address_match = (await _address_matches_by_rows(
            cur,
            parser_type,
            [row_key],
        )).get(row_key)
    source_count = int(projection[1] or 0)
    conflict = bool(projection[2])
    pending = bool(str(projection[3] or ""))
    task_state_value = str(projection[4] or "")
    task = _task_record(
        parser_type,
        row_key,
        values,
        source_count,
        conflict,
        pending,
        task_state_value,
        registration_link=registration_link,
        review_flow=review_flow,
        address_match=address_match,
    )
    return {
        "task_update": {
            key: task[key]
            for key in (
                "row_key", "summary", "community", "inspector", "state",
                "needs_review", "review_stage", "review_flow", "source_count",
                "conflict", "pending_sync", "sync_state", "priority",
                "registration_link", "address_match",
            )
        },
        "registration_link": registration_link,
        "review_flow": review_flow,
    }


@router.post("/{parser_type}/{row_key}/registration/confirm")
async def manually_confirm_registration(
    parser_type: str,
    row_key: str,
    data: RegistrationManualConfirm,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """高权限人工确认登记，仅限居住证有效但自动匹配需复核的任务。"""
    if parser_type not in REGISTRATION_TASK_TYPES:
        raise HTTPException(400, "该业务不支持登记闭环")
    if not can_confirm_registration(user):
        raise HTTPException(403, "只有基础管控及以上岗位可以人工确认登记")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    config = await load_residence_config(conn)
    if not config.session_ready:
        raise HTTPException(409, "居住证平台配置尚未就绪，不能人工确认")
    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT link.status,link.reason_code,link.source_id,projection.identity_hmac,"
            "projection.community,projection.source_count "
            "FROM _task_registration_links AS link "
            "JOIN _online_source_projection AS projection "
            "ON projection.parser_type=link.parser_type AND projection.row_key=link.row_key "
            "WHERE link.parser_type=%s AND link.row_key=%s",
            (parser_type, row_key),
        )
        link = await cur.fetchone()
        if not link or str(link[0] or "") != "review_required" or str(link[1] or "") not in {
            "address_mismatch", "address_ambiguous"
        }:
            raise HTTPException(409, "该任务当前不满足人工确认条件")
        if int(link[5] or 0) != 1:
            raise HTTPException(409, "任务来源不唯一，不能人工确认")
        community_values = context.get("community_values")
        if community_values is not None and str(link[4] or "") not in set(community_values):
            raise HTTPException(403, "无权复核该社区任务")
        expected_hmac = str(link[3] or "")
        await cur.execute(
            "SELECT values_json,revision FROM _online_source_rows "
            "WHERE id=%s AND parser_type=%s AND row_key=%s",
            (int(link[2]), parser_type, row_key),
        )
        source_row = await cur.fetchone()
        if not source_row:
            raise HTTPException(404, "任务来源已不存在")
        current_source_values = json_value(source_row[0], {})
        current_revision = int(source_row[1])
    await conn.rollback()
    if data.expected_revision != current_revision:
        raise HTTPException(409, "任务来源已变化，请刷新后重试")
    try:
        target = await _load_current_target(parser_type, row_key, expected_hmac)
        state, _address, registration_code = await _lookup_registration_address_target(config, target)
    except Exception as exc:
        raise HTTPException(502, "居住证平台查询失败，不能人工确认") from exc
    if state != "registered" or registration_code == "1":
        raise HTTPException(409, "居住证登记无效或已注销，不能人工确认")

    async def confirm_prepare(*, cur, source, current_values, changes):
        del changes
        workflow = TASK_WORKFLOWS[parser_type]
        if str(current_values.get(workflow.result_field) or "").strip() != "待登记":
            raise HTTPException(409, "任务当前已不是待登记，请刷新后重试")
        await cur.execute(
            "SELECT status,reason_code,source_id,property_id "
            "FROM _task_registration_links WHERE parser_type=%s AND row_key=%s FOR UPDATE",
            (parser_type, row_key),
        )
        current_link = await cur.fetchone()
        if (
            not current_link
            or str(current_link[0] or "") != "review_required"
            or str(current_link[1] or "") not in {"address_mismatch", "address_ambiguous"}
            or int(current_link[2] or 0) != int(source["id"])
            or current_link[3] is None
        ):
            raise HTTPException(409, "登记复核状态已变化，请刷新后重试")
        return {}

    async def confirm_callback(*, cur, source, before, after, row_key_before, row_key_after, revision):
        del source, before, after, revision
        await cur.execute(
            "UPDATE _task_registration_links SET status='confirmation_pending',reason_code='',"
            "confirmed_by=%s,manual_confirmed_at=UTC_TIMESTAMP(),confirmed_at=NULL,"
            "manual_reason=%s,manual_note=%s "
            "WHERE parser_type=%s AND row_key=%s AND source_id=%s AND status='review_required'",
            (int(user.get("id")), data.reason, data.note[:500], parser_type, row_key_after, int(link[2])),
        )
        if cur.rowcount != 1:
            raise HTTPException(409, "登记复核状态已变化，请刷新后重试")
        await record_registration_event(
            cur, parser_type=parser_type, row_key=row_key_after,
            source_id=int(link[2]), event_type="manual_registration_confirmed",
            reason_code=data.reason, actor_user_id=int(user.get("id")),
        )

    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=int(link[2]),
        changes={TASK_WORKFLOWS[parser_type].result_field: "已登记"},
        base_values={TASK_WORKFLOWS[parser_type].result_field: str(current_source_values.get(TASK_WORKFLOWS[parser_type].result_field) or "")},
        request=request, user=user, conn=conn, explicit_text_edit=False,
        registration_mode=True, audit_action="manual_registration",
        transaction_prepare=confirm_prepare,
        transaction_callback=confirm_callback,
    )
    await record_admin_audit(
        user, "mobile_tasks.registration.manual_confirm",
        target_type="mobile_task", target_name=f"{parser_type}:{row_key}",
        detail={"reason": data.reason}, **request_audit_fields(request),
    )
    return result


@router.patch("/{parser_type}/source-rows/{source_id}")
async def update_mobile_task(
    parser_type: str,
    source_id: int,
    data: TaskBatchUpdate,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    user = _require_task_edit_user(user)
    analysis_fields = set(TASK_WORKFLOWS[parser_type].analysis_fields)
    if any(field in analysis_fields for field in data.changes):
        raise HTTPException(403, "请从研判页面修改研判内容")
    context = await _flow_context(conn, user)
    async with conn.cursor() as cur:
        await _validate_assignment(cur, context, data.changes)
    transaction_prepare, transaction_callback, registration_mode = (
        _registration_update_hooks(parser_type, data, user)
    )
    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=source_id,
        changes=data.changes,
        base_values=data.base_values,
        expected_revision=data.expected_revision,
        request=request,
        user=user,
        conn=conn,
        explicit_text_edit=True,
        registration_mode=registration_mode,
        transaction_prepare=transaction_prepare,
        transaction_callback=transaction_callback,
    )
    if registration_mode:
        wake_residence_lookup_scheduler()
    return {
        **result,
        **await _mobile_task_save_context(conn, parser_type, str(result["row_key"])),
    }


@router.post("/{parser_type}/source-rows/{source_id}/claim")
async def claim_mobile_task(
    parser_type: str,
    source_id: int,
    data: TaskBatchUpdate,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """Let a group member claim an unassigned task while saving their edit."""
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if context.get("admin_mode") or context.get("position") != "组员":
        raise HTTPException(403, "只有组员可以自主领取未分配任务")
    if not data.changes:
        raise HTTPException(400, "请先填写需要保存的核查内容")
    if "核查人" in data.changes:
        raise HTTPException(400, "领取人由当前登录账号自动确定")
    analysis_fields = set(TASK_WORKFLOWS[parser_type].analysis_fields)
    if any(field in analysis_fields for field in data.changes):
        raise HTTPException(403, "请从研判页面修改研判内容")

    inspector = str(context["name"] or "").strip()
    transaction_prepare, transaction_callback, registration_mode = (
        _registration_update_hooks(parser_type, data, user)
    )

    def validate_unassigned(current: dict) -> None:
        if str(current.get("核查人") or "").strip():
            raise HTTPException(409, "该任务已被领取，请刷新后重试")

    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=source_id,
        changes={"核查人": inspector, **data.changes},
        base_values={**data.base_values, "核查人": ""},
        expected_revision=data.expected_revision,
        request=request,
        user=user,
        conn=conn,
        explicit_text_edit=True,
        current_values_validator=validate_unassigned,
        registration_mode=registration_mode,
        transaction_prepare=transaction_prepare,
        transaction_callback=transaction_callback,
    )
    if registration_mode:
        wake_residence_lookup_scheduler()
    await record_admin_audit(
        user,
        "mobile_tasks.self_claim",
        target_type="mobile_task",
        target_name=f"{parser_type}:{source_id}",
        detail={"columns": list(data.changes)},
        **request_audit_fields(request),
    )
    return {
        **result,
        **await _mobile_task_save_context(conn, parser_type, str(result["row_key"])),
        "message": "已领取任务并保存到本地任务池",
    }


@router.get("/{parser_type}/internal-transfer-options")
async def get_internal_transfer_options(
    parser_type: str,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    _require_unverifiable_reviewer(user)
    async with conn.cursor() as cur:
        await cur.execute(
            """
            SELECT community.name,member.name
            FROM _communities community
            JOIN _departments department
              ON department.community_id=community.id
             AND department.department_type='community'
             AND department.is_active=1
            LEFT JOIN _grid_member_department_links link
              ON link.department_id=department.id
            LEFT JOIN _grid_members member
              ON member.id=link.member_id
             AND member.position='组长'
             AND member.status='在岗'
            WHERE community.is_active=1
            ORDER BY community.name,member.name
            """
        )
        grouped: dict[str, list[str]] = {}
        for community, leader in await cur.fetchall():
            name = str(community or "").strip()
            if not name:
                continue
            grouped.setdefault(name, [])
            leader_name = str(leader or "").strip()
            if leader_name and leader_name not in grouped[name]:
                grouped[name].append(leader_name)
    return {
        "data": [
            {"community": community, "leaders": leaders}
            for community, leaders in grouped.items()
        ],
    }


@router.post("/{parser_type}/source-rows/{source_id}/internal-transfer")
async def transfer_mobile_task_internally(
    parser_type: str,
    source_id: int,
    data: InternalTransferRequest,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    scoped_user = _require_unverifiable_reviewer(user)
    target_community = data.target_community.strip()
    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT row_key,row_hash FROM _online_source_rows source "
            "WHERE source.id=%s AND source.parser_type=%s AND source.archived_at IS NULL "
            f"{active_source_sql_filter(parser_type)}",
            (source_id, parser_type),
        )
        source_row = await cur.fetchone()
        if not source_row:
            raise HTTPException(404, "本地任务来源不存在")
        if str(source_row[0] or "") != data.expected_row_key:
            raise HTTPException(409, "任务主键已变化，请刷新后重试")
        if str(source_row[1] or "") != data.expected_row_hash:
            raise HTTPException(409, "任务已被其他人更新，请刷新后重试")
        await cur.execute(
            """
            SELECT member.name
            FROM _communities community
            JOIN _departments department
              ON department.community_id=community.id
             AND department.department_type='community'
             AND department.is_active=1
            JOIN _grid_member_department_links link
              ON link.department_id=department.id
            JOIN _grid_members member
              ON member.id=link.member_id
             AND member.position='组长'
             AND member.status='在岗'
            WHERE community.is_active=1 AND community.name=%s
            ORDER BY member.name
            """,
            (target_community,),
        )
        leaders = list(dict.fromkeys(
            str(row[0] or "").strip() for row in await cur.fetchall()
            if str(row[0] or "").strip()
        ))
    if not leaders:
        raise HTTPException(409, "目标社区没有在岗组长，请先维护人员配置")
    if len(leaders) == 1:
        target_leader = leaders[0]
    else:
        target_leader = data.target_leader.strip()
        if not target_leader:
            raise HTTPException(422, "目标社区有多名在岗组长，请明确选择接手组长")
        if target_leader not in leaders:
            raise HTTPException(400, "所选人员不是目标社区的在岗组长")

    workflow = TASK_WORKFLOWS[parser_type]
    community_field = get_parser(parser_type).COMMUNITY_COLUMN

    def validate_transfer(current: dict) -> None:
        result = str(current.get(workflow.result_field) or "").strip()
        if result != "移交（所内）":
            raise HTTPException(409, "只有核查结果为“移交（所内）”的任务可以转交")

    async def transfer_callback(
        *, cur, source, before, after, row_key_before, row_key_after, revision,
    ):
        del after, row_key_before
        try:
            await record_internal_transfer(
                cur,
                parser_type=parser_type,
                row_key=row_key_after,
                source_id=int(source["id"]),
                before=before,
                target_community=target_community,
                target_leader=target_leader,
                operator_user_id=int(scoped_user.get("id")) if scoped_user.get("id") else None,
                source_revision=revision,
                community_field=community_field,
            )
        except LookupError as exc:
            raise HTTPException(409, str(exc)) from exc

    result = await queue_source_fields(
        parser_type=parser_type,
        source_id=source_id,
        changes={community_field: target_community, "核查人": target_leader},
        base_values={},
        expected_revision=data.expected_revision,
        request=request,
        user=scoped_user,
        conn=conn,
        explicit_text_edit=True,
        current_values_validator=validate_transfer,
        transaction_callback=transfer_callback,
        audit_action="internal_transfer_local",
    )
    await record_admin_audit(
        scoped_user,
        "mobile_tasks.internal_transfer",
        target_type="mobile_task",
        target_name=f"{parser_type}:{source_id}",
        detail={"target_community": target_community, "target_leader": target_leader},
        **request_audit_fields(request),
    )
    return {**result, "target_community": target_community, "target_leader": target_leader}


@router.post("/{parser_type}/source-rows/{source_id}/resolve-sync-conflict")
async def resolve_mobile_task_sync_conflict(
    parser_type: str,
    source_id: int,
    data: SyncConflictResolution,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入手机任务工作台")
    fields = list(dict.fromkeys(str(field).strip() for field in data.fields if str(field).strip()))
    if not fields:
        raise HTTPException(400, "请选择需要处理的冲突字段")

    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT row_key FROM _online_source_rows WHERE id=%s AND parser_type=%s",
            (source_id, parser_type),
        )
        row = await cur.fetchone()
        if not row:
            raise HTTPException(404, "来源行不存在")
        row_key = str(row[0])

    analysis_fields = set(TASK_WORKFLOWS[parser_type].analysis_fields)
    analysis_mode = any(field in analysis_fields for field in fields)
    if analysis_mode and any(field not in analysis_fields for field in fields):
        raise HTTPException(400, "核查字段和研判字段的冲突请分别处理")
    scoped_user = (
        _require_analysis_user(user)
        if analysis_mode else _require_task_edit_user(user)
    )
    detail = await _mobile_task_detail_data(
        parser_type,
        row_key,
        scoped_user,
        conn,
        analysis_mode=analysis_mode,
        include_photo_requests=False,
    )
    source = next((item for item in detail["sources"] if item["id"] == source_id), None)
    if not source:
        raise HTTPException(404, "来源行不存在或不属于当前社区")
    conflict_fields = {
        item["field"] for item in source.get("sync_fields", [])
        if item.get("status") == "conflict"
    }
    if any(field not in conflict_fields for field in fields):
        raise HTTPException(409, "所选字段已不处于同步冲突状态，请刷新后重试")
    if any(field not in source.get("editable_fields", []) for field in fields):
        raise HTTPException(403, "当前账号无权处理所选字段")

    try:
        result = await resolve_source_conflict(
            conn,
            source_id=source_id,
            choice=data.choice,
            fields=fields,
        )
    except LookupError as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    await record_admin_audit(
        scoped_user,
        "online.writeback.resolve_conflict",
        target_type="online_source_row",
        target_name=f"{parser_type}:{source_id}",
        detail={"choice": data.choice, "columns": fields},
        **request_audit_fields(request),
    )
    if data.choice == "platform":
        launch_local_change_processing(source_id)
    return {
        "message": (
            (
                "已采用平台值并保存到本地任务池"
                if local_data_source_enabled()
                else "已采用平台值，正在重新同步腾讯表格"
            )
            if data.choice == "platform"
            else (
                "已采用本地任务值"
                if local_data_source_enabled()
                else "已采用腾讯值"
            )
        ),
        **result,
    }


@router.post("/{parser_type}/bulk-assign")
async def bulk_assign_mobile_tasks(
    parser_type: str,
    data: BulkAssignmentRequest,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """Assign one bounded local task chunk with one incremental projection pass."""
    started_at = time.monotonic()
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if not _can_assign_tasks(context):
        raise HTTPException(403, "只有组长及有权管理任务的上级岗位可以批量分配核查人")

    row_keys = list(dict.fromkeys(str(value).strip() for value in data.row_keys if str(value).strip()))
    inspector = str(data.inspector or "").strip()
    if not row_keys:
        raise HTTPException(400, "请选择任务")
    if data.mode == "single" and not inspector:
        raise HTTPException(400, "请选择在岗组员")
    balanced_total = data.balanced_total or len(row_keys)
    if data.mode == "balanced" and (
        data.balanced_offset + len(row_keys) > balanced_total
    ):
        raise HTTPException(400, "平均分配分块范围无效，请重新发起分配")

    projection_by_key: dict[str, dict] = {}
    source_rows_by_key: dict[str, list[dict]] = {}
    skipped: list[dict[str, str]] = []
    async with conn.cursor() as cur:
        if not local_data_source_enabled() and not await _writeback_enabled(cur):
            raise HTTPException(503, "在线回写已由超级管理员暂停")
        scope_where, scope_params = _scope_where(
            context,
            "all" if context.get("admin_mode") else "community",
        )
        key_placeholders = ", ".join(["%s"] * len(row_keys))
        await cur.execute(
            f"""
            SELECT projection.row_key, projection.community, projection.inspector,
                   projection.task_state, projection.conflict, projection.source_count,
                   projection.address_match_status,
                   projection.small_community_id,
                   projection.small_community_name
            FROM _online_source_projection AS projection
            WHERE projection.parser_type=%s
              AND projection.row_key IN ({key_placeholders})
              AND {scope_where}
            """,
            [parser_type, *row_keys, *scope_params],
        )
        for row in await cur.fetchall():
            projection_by_key[str(row[0])] = {
                "community": str(row[1] or "").strip(),
                "inspector": str(row[2] or "").strip(),
                "state": str(row[3] or "").strip(),
                "conflict": bool(row[4]),
                "source_count": int(row[5] or 0),
                "address_match_status": str(row[6] or "unmatched"),
                "small_community_id": int(row[7]) if row[7] is not None else None,
                "small_community_name": str(row[8] or ""),
            }
        missing_scope = [key for key in row_keys if key not in projection_by_key]
        if missing_scope:
            raise HTTPException(403, "部分任务不在当前账号可操作范围内，请刷新后重试")

        assignment_context = await inspector_option_context(
            cur,
            user,
            assignment_only=True,
        )
        community_aliases = assignment_context["community_aliases"]
        inspectors_by_community = assignment_context["inspectors_by_community"]
        formal_communities = {
            community_aliases.get(item["community"], "")
            for item in projection_by_key.values()
        }
        if not formal_communities or "" in formal_communities:
            raise HTTPException(403, "部分任务社区不在当前账号可分配范围内")
        if len(formal_communities) != 1:
            raise HTTPException(400, "批量分配必须一次只选择同一社区的任务")
        formal_community = next(iter(formal_communities))
        available_inspectors = list(
            inspectors_by_community.get(formal_community) or []
        )
        if not available_inspectors:
            raise HTTPException(400, "该社区当前没有在岗组员可分配")
        if data.mode == "single" and inspector not in available_inspectors:
            raise HTTPException(400, "只能分配给任务所属社区的在岗组员")

        await cur.execute(
            f"""
            SELECT source.id, source.row_key, source.revision,
                   source.physical_row, source.row_hash, source.values_json,
                   source.sheet_id
            FROM _online_source_rows AS source
            WHERE source.parser_type=%s
              AND source.row_key IN ({key_placeholders})
              AND source.archived_at IS NULL
              {active_source_sql_filter(parser_type)}
            ORDER BY source.row_key, source.id
            """,
            [parser_type, *row_keys],
        )
        for source_id, row_key, revision, physical_row, row_hash, values_json, sheet_id in await cur.fetchall():
            source_rows_by_key.setdefault(str(row_key), []).append({
                "id": int(source_id),
                "row_key": str(row_key),
                "revision": int(revision),
                "physical_row": int(physical_row),
                "row_hash": str(row_hash or ""),
                "values": json_value(values_json, {}),
                "sheet_id": str(sheet_id or ""),
            })
        if local_data_source_enabled():
            # Projection rows may still carry the old Tencent source count or
            # conflict bit until the first local rebuild.  Assignment must
            # use the active local rows just loaded above.
            for row_key, item in projection_by_key.items():
                active_count = len(source_rows_by_key.get(row_key) or [])
                item["source_count"] = active_count
                item["conflict"] = active_count > 1

    eligible_keys: list[str] = []
    for row_key in row_keys:
        item = projection_by_key[row_key]
        if item["inspector"]:
            skipped.append({"row_key": row_key, "reason": "已有核查人"})
            continue
        if item["state"] == "completed":
            skipped.append({"row_key": row_key, "reason": "任务已完成"})
            continue
        if item["conflict"]:
            skipped.append({"row_key": row_key, "reason": "来源存在冲突"})
            continue
        if item["source_count"] != 1:
            skipped.append({"row_key": row_key, "reason": "存在重复本地来源，请先处理来源异常"})
            continue
        if item["address_match_status"] not in {"confirmed", "suggested"} or not item["small_community_id"]:
            skipped.append({"row_key": row_key, "reason": "小区归属未形成唯一可靠结果，请先处理地址匹配"})
            continue
        if not source_rows_by_key.get(row_key):
            skipped.append({"row_key": row_key, "reason": "找不到来源行"})
            continue
        eligible_keys.append(row_key)

    if data.mode == "balanced":
        assignment_plan, assignment_counts = _balanced_assignment_plan(
            row_keys,
            available_inspectors,
            total_count=balanced_total,
            start_index=data.balanced_offset,
        )
    else:
        assignment_plan = {row_key: inspector for row_key in eligible_keys}
        assignment_counts = {
            inspector: len(eligible_keys),
        }

    update_count = 0
    failures: list[dict[str, str]] = []
    successful_keys: list[str] = []
    successful_assignment_counts = {name: 0 for name in assignment_counts}
    if local_data_source_enabled():
        await conn.begin()
        try:
            async with conn.cursor() as cur:
                for row_key in eligible_keys:
                    assigned_inspector = assignment_plan.get(row_key, "")
                    if not assigned_inspector:
                        skipped.append({"row_key": row_key, "reason": "没有生成分配方案"})
                        continue
                    sources = source_rows_by_key.get(row_key) or []
                    if len(sources) != 1:
                        skipped.append({"row_key": row_key, "reason": "本地来源不唯一"})
                        continue
                    source = sources[0]
                    await cur.execute("SAVEPOINT bulk_assign_task")
                    try:
                        await apply_local_system_changes(
                            cur,
                            source={
                                **source,
                                "spreadsheet_id": 0,
                                "spreadsheet": {"parser_type": parser_type},
                            },
                            changes={"核查人": assigned_inspector},
                            user=user,
                            action="bulk_assign_local",
                            rebuild=False,
                        )
                    except (ValueError, LookupError):
                        await cur.execute("ROLLBACK TO SAVEPOINT bulk_assign_task")
                        failures.append({"row_key": row_key, "reason": "任务已变化，请刷新后重试"})
                    else:
                        await capture_first_assignment(
                            cur,
                            parser_type=parser_type,
                            row_key=row_key,
                            community=projection_by_key[row_key]["community"],
                            inspector=assigned_inspector,
                            actor_user_id=int(user.get("id")) if user.get("id") else None,
                            source="bulk_assignment",
                        )
                        update_count += 1
                        successful_keys.append(row_key)
                        successful_assignment_counts[assigned_inspector] = (
                            successful_assignment_counts.get(assigned_inspector, 0) + 1
                        )
                    finally:
                        await cur.execute("RELEASE SAVEPOINT bulk_assign_task")
            await conn.commit()
        except Exception:
            await conn.rollback()
            raise
    else:
        for row_key in eligible_keys:
            assigned_inspector = assignment_plan.get(row_key, "")
            if not assigned_inspector:
                skipped.append({"row_key": row_key, "reason": "没有生成分配方案"})
                continue
            sources = source_rows_by_key.get(row_key) or []
            if not sources:
                skipped.append({"row_key": row_key, "reason": "找不到来源行"})
                continue
            task_ok = True
            for source in sources:
                try:
                    await queue_source_fields(
                        parser_type=parser_type,
                        source_id=source["id"],
                        changes={"核查人": assigned_inspector},
                        expected_revision=source["revision"],
                        request=request,
                        user=user,
                        conn=conn,
                        explicit_text_edit=True,
                    )
                except HTTPException as exc:
                    task_ok = False
                    reason = {
                        400: "数据校验未通过",
                        403: "没有该任务的编辑权限",
                        409: "任务已变化，请刷新后重试",
                        502: "腾讯回写校验失败",
                    }.get(exc.status_code, "保存失败")
                    failures.append({"row_key": row_key, "reason": reason})
                    break
            if task_ok:
                update_count += 1
                successful_keys.append(row_key)
                successful_assignment_counts[assigned_inspector] = (
                    successful_assignment_counts.get(assigned_inspector, 0) + 1
                )

    result = _bulk_assignment_result(
        updated=update_count,
        skipped=skipped,
        failures=failures,
        inspector=inspector,
        mode=data.mode,
        assignment_counts=successful_assignment_counts,
        duration_ms=max(0, int((time.monotonic() - started_at) * 1000)),
    )
    await record_admin_audit(
        user,
        "mobile_tasks.bulk_assign",
        target_type="mobile_task",
        target_name=parser_type,
        detail={
            "task_count": len(row_keys),
            "updated_count": result["updated"],
            "skipped_count": result["skipped"],
            "failed_count": result["failed"],
            "inspector": inspector if data.mode == "single" else "",
            "allocation_mode": data.mode,
            "assignment_counts": successful_assignment_counts,
        },
        **request_audit_fields(request),
    )
    logger.info(
        "mobile_task_bulk_assignment parser=%s requested=%d eligible=%d "
        "updated=%d skipped=%d failed=%d duration_ms=%d",
        parser_type,
        len(row_keys),
        len(eligible_keys),
        result["updated"],
        result["skipped"],
        result["failed"],
        max(0, int((time.monotonic() - started_at) * 1000)),
    )
    return result


@router.post("/{parser_type}/assignments/cancel")
async def cancel_mobile_task_assignments(
    parser_type: str,
    data: CancelAssignmentsRequest,
    request: Request,
    user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """撤销当前业务范围内尚未完成任务的核查人分配。

    该操作只清除“核查人”字段，不修改核查结果、研判、反馈或任务历史；
    每条变更仍通过本地事务写入审计，便于追溯误分配的恢复过程。
    """
    if parser_type not in TASK_WORKFLOWS:
        raise HTTPException(400, "该业务尚未接入任务工作台")
    user = _require_task_edit_user(user)
    context = await _flow_context(conn, user)
    if not _can_assign_tasks(context):
        raise HTTPException(403, "只有组长及有权管理任务的上级岗位可以撤销批量分配")
    if not local_data_source_enabled():
        raise HTTPException(410, "本地数据源未启用，无法执行本地分配恢复")

    requested_community = str(data.community or "").strip()
    scope_where, scope_params = _scope_where(
        context,
        "all" if context.get("admin_mode") else "community",
    )
    if requested_community:
        aliases = await community_names_for_scopes(conn, [requested_community])
        formal = aliases[0] if aliases else requested_community
        scope_where += " AND projection.community=%s"
        scope_params.append(formal)

    changed = 0
    skipped: list[dict[str, str]] = []
    await conn.begin()
    try:
        async with conn.cursor() as cur:
            await cur.execute(
                f"""
                SELECT source.id, source.row_key, source.revision,
                       source.physical_row, source.values_json,
                       projection.task_state, projection.inspector,
                       projection.source_count, projection.conflict
                FROM _online_source_projection projection
                JOIN _online_source_rows source
                  ON source.parser_type=projection.parser_type
                 AND source.row_key=projection.row_key
                 AND source.archived_at IS NULL
                 AND source.spreadsheet_id=0
                WHERE projection.parser_type=%s
                  AND {scope_where}
                  AND TRIM(COALESCE(projection.inspector,''))<>''
                  AND projection.task_state<>'completed'
                ORDER BY source.id
                LIMIT %s
                """,
                [parser_type, *scope_params, MAX_BULK_ASSIGNMENT_TASKS * 20],
            )
            rows = await cur.fetchall()
            for source_id, row_key, revision, physical_row, values_json, task_state, inspector, source_count, conflict in rows:
                if int(source_count or 0) != 1 or bool(conflict):
                    skipped.append({"row_key": str(row_key), "reason": "来源异常，未自动撤销"})
                    continue
                values = json_value(values_json, {})
                if not str(values.get("核查人") or "").strip():
                    continue
                try:
                    await apply_local_system_changes(
                        cur,
                        source={
                            "id": int(source_id),
                            "revision": int(revision),
                            "physical_row": int(physical_row),
                            "row_key": str(row_key),
                            "values": values,
                            "spreadsheet_id": 0,
                            "spreadsheet": {"parser_type": parser_type},
                        },
                        changes={"核查人": ""},
                        user=user,
                        action="bulk_unassign_local",
                        rebuild=False,
                    )
                    changed += 1
                except (ValueError, LookupError) as exc:
                    skipped.append({"row_key": str(row_key), "reason": str(exc) or "任务已变化，请刷新后重试"})
        await conn.commit()
    except Exception:
        await conn.rollback()
        raise

    await record_admin_audit(
        user,
        "mobile_tasks.bulk_unassign",
        target_type="mobile_task",
        target_name=parser_type,
        detail={
            "community": requested_community,
            "updated_count": changed,
            "skipped_count": len(skipped),
        },
        **request_audit_fields(request),
    )
    return {"updated": changed, "skipped": len(skipped), "details": skipped}
