"""Flat, signed workbooks for offline property-to-community annotation."""

from __future__ import annotations

import hashlib
import hmac
import json
from collections import Counter
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.exceptions import InvalidFileException

from config import settings

FORMAT = "binhu-property-annotation-v2"
MAX_ROWS = 10000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SNAPSHOT_FIELDS = (
    "id", "version", "community_id", "community_name", "natural_address",
    "normalized_address", "street", "building", "room", "status",
    "updated_at", "small_community_id", "address_match_status",
    "address_match_confirmed_by", "address_match_confirmed_at", "household_status",
)
LABEL_FIELDS = ("decision", "annotated_small_community_id", "annotation_reason")
HEADERS = (*SNAPSHOT_FIELDS, "address_match_score", "address_match_reason",
           "snapshot_token", *LABEL_FIELDS)
ENTRY_FIELDS = ("id", "name", "community_id", "community_name", "detail_address",
                "aliases", "address_type")


def signature(kind: str, data: dict, user_id: int) -> str:
    body = json.dumps([FORMAT, kind, int(user_id), data], ensure_ascii=False,
                      sort_keys=True, separators=(",", ":"), default=str)
    return hmac.new(settings.registry_hmac_key.encode(), body.encode(), hashlib.sha256).hexdigest()


def property_snapshot(row: dict) -> dict:
    return {key: "" if row.get(key) is None else str(row[key]) for key in SNAPSHOT_FIELDS}


def property_token(row: dict, user_id: int) -> str:
    return signature("property", property_snapshot(row), user_id)


def entry_token(entry: dict, user_id: int) -> str:
    return signature("entry", {key: entry.get(key) for key in ENTRY_FIELDS}, user_id)


def valid_token(actual: str, expected: str) -> bool:
    return isinstance(actual, str) and actual.isascii() and hmac.compare_digest(actual, expected)


def _sheet(workbook, name, headers, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(list(headers))
    for index, row in enumerate(rows, 2):
        sheet.append([value if value is not None else "" for value in row])
        # Store all strings explicitly as text, including strings starting '='.
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                sheet.cell(index, column).data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="28645B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 42
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            60, max(20, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return sheet


def build_annotation_workbook(properties: list[dict], entries: list[dict], user_id: int) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = _sheet(workbook, "房屋标注", HEADERS, [
        [*property_snapshot(row).values(), row.get("address_match_score", 0),
         row.get("address_match_reason", ""), property_token(row, user_id), "", "", ""]
        for row in properties
    ])
    decision_column = sheet.cell(1, HEADERS.index("decision") + 1).column_letter
    if properties:
        validation = DataValidation(type="list", formula1='"match,review,skip"', allow_blank=True)
        validation.errorTitle = "标注决定无效"
        validation.error = "请选择 match、review 或 skip"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{decision_column}2:{decision_column}{len(properties) + 1}")
        for row in sheet.iter_rows(min_row=2):
            for cell in row[-len(LABEL_FIELDS):]:
                cell.fill = PatternFill("solid", fgColor="E7F1FB")
    _sheet(workbook, "小区地址库", (*ENTRY_FIELDS, "snapshot_token"), [
        [*(json.dumps(entry.get(key), ensure_ascii=False) if key == "aliases" else entry.get(key)
           for key in ENTRY_FIELDS), entry_token(entry, user_id)] for entry in entries
    ])
    entry_ids = {entry["id"] for entry in entries}
    candidate_rows = []
    for row in properties:
        for rank, candidate in enumerate(row.get("address_match_candidates") or [], 1):
            if not isinstance(candidate, dict) or candidate.get("entry_id") not in entry_ids:
                continue
            candidate_rows.append([row["id"], rank, candidate.get("entry_id"),
                                   candidate.get("score"), candidate.get("method"), candidate.get("reason")])
    _sheet(workbook, "候选小区", ("property_id", "rank", "small_community_id", "score", "method", "reason"), candidate_rows)
    _sheet(workbook, "标注说明", ("key", "value"), [
        ["format", FORMAT], ["source", "本地 MySQL 房屋档案与小区管理；按导出账号权限、当前筛选和排序生成"],
        ["editable_columns", ", ".join(LABEL_FIELDS)],
        ["match", "选择唯一小区：decision=match，填写小区地址库的 id 和 annotation_reason（依据）；回导后由管理员确认"],
        ["review", "无法确定、冲突或小区库缺失：decision=review，填写原因；保持房屋当前归属"],
        ["skip", "不处理：decision=skip 或留空；不会修改档案"],
        ["ids", "id 为稳定标识；严禁按物理行号或名称推断关联。选择的小区必须与房屋所属社区一致"],
        ["read_only", "只填写三个标注列。保留其余原始列、工作表、签名和版本；不得修改原始地址或社区"],
        ["scope", "小区地址库包含当前账号可查看社区的启用小区，aliases 为 JSON 数组；不是全平台无权限数据"],
        ["household_status", "注销状态使用 household_status；status 仅为历史兼容原始列，不代表户号表注销状态，不要改动"],
        ["conflicts", "房屋已变化、小区已停用、跨社区、重复行或签名失效均拒绝应用；重新导出后标注"],
        ["review_boundary", "Agent 结果仅为建议；已有人工确认的变更在预览中明确显示，只有管理员确认才应用"],
        ["privacy", "仅含地址匹配所需资料；不含人员、房东姓名、身份证、电话、核查或走访正文。按授权范围交给外部 Agent"],
        ["return_account", "由原导出账号回导与确认；其他账号需重新导出"],
    ])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def parse_annotation_workbook(content: bytes, user_id: int) -> tuple[list[dict], dict[int, str]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("标注文件超过 10MB，请按社区拆分导出")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 200 or sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("工作簿解压体积超限，请按社区拆分")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException, ValueError) as exc:
        raise ValueError("无法读取 XLSX 标注工作簿") from exc
    try:
        if not {"房屋标注", "小区地址库", "标注说明"}.issubset(workbook.sheetnames):
            raise ValueError("请使用平台导出的小区标注工作簿，保留工作表")
        instructions_sheet = workbook["标注说明"]
        if instructions_sheet.max_row > 100:
            raise ValueError("标注说明已变化，请重新导出")
        instructions = dict(instructions_sheet.iter_rows(min_row=2, max_row=100, max_col=2, values_only=True))
        if instructions.get("format") != FORMAT:
            raise ValueError("工作簿格式版本不支持，请重新导出")
        sheet = workbook["房屋标注"]
        if sheet.max_row > MAX_ROWS + 1:
            raise ValueError("工作簿超过 10000 行，请按社区拆分")
        headers = next(sheet.iter_rows(max_row=1, values_only=True), ())
        if tuple(headers) != HEADERS:
            raise ValueError("房屋标注表头已变化，请保留原表头和原始列")
        rows = []
        for index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(HEADERS)), 2):
            if len(rows) >= MAX_ROWS:
                raise ValueError("工作簿超过 10000 行，请按社区拆分")
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            row = dict(zip(HEADERS, values))
            row["xlsx_row"] = index
            row["formula"] = any(cell.data_type == "f" for cell in cells)
            rows.append(row)
        counts = Counter(str(row.get("id") or "") for row in rows)
        for row in rows:
            row["duplicate"] = counts[str(row.get("id") or "")] > 1
        catalog = workbook["小区地址库"]
        if catalog.max_row > MAX_ROWS + 1:
            raise ValueError("小区地址库超过 10000 行")
        if tuple(next(catalog.iter_rows(max_row=1, values_only=True), ())) != (*ENTRY_FIELDS, "snapshot_token"):
            raise ValueError("小区地址库表头已变化，请重新导出")
        entry_tokens = {}
        for index, cells in enumerate(catalog.iter_rows(min_row=2, max_col=len(ENTRY_FIELDS) + 1), 2):
            if index > MAX_ROWS + 1:
                raise ValueError("小区地址库超过 10000 行")
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("小区地址库不能包含公式")
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            try:
                entry_id = positive_id(values[0])
            except ValueError as exc:
                raise ValueError("小区地址库 id 无效") from exc
            if entry_id in entry_tokens:
                raise ValueError("小区地址库存在重复 id")
            entry = dict(zip(ENTRY_FIELDS, values))
            entry["id"] = entry_id
            for key in ("name", "community_name", "detail_address", "address_type"):
                entry[key] = str(entry[key] or "")
            try:
                entry["aliases"] = json.loads(entry["aliases"] or "[]")
                entry["community_id"] = positive_id(entry["community_id"])
            except (ValueError, TypeError) as exc:
                raise ValueError("小区地址库原始资料已变化，请重新导出") from exc
            if not valid_token(str(values[-1] or ""), entry_token(entry, user_id)):
                raise ValueError("小区地址库原始资料已变化或非原导出账号，请重新导出")
            entry_tokens[entry_id] = str(values[-1] or "")
        return rows, entry_tokens
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException) as exc:
        raise ValueError("无法读取 XLSX 标注工作簿") from exc
    finally:
        workbook.close()


def positive_id(value) -> int:
    text = str(value or "")
    if not text.isascii() or not text.isdecimal() or int(text) <= 0:
        raise ValueError("ID 必须为正整数")
    return int(text)
