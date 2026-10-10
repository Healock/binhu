from __future__ import annotations

import os
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

os.environ.setdefault("MYSQL_PASSWORD", "test-password")
os.environ.setdefault("ENCRYPTION_KEY", "test-encryption-key")

from routers import registry
from services import property_annotation_xlsx as xlsx


USER = {"id": 7}
PROPERTY = {"id": 42, "version": 3, "community_id": 8, "community_name": "合成社区",
            "natural_address": "合成路1号", "normalized_address": "合成路1号",
            "street": "合成街道", "building": "1", "room": "101", "status": "active",
            "updated_at": "2026-10-10T01:00:00", "small_community_id": None,
            "household_status": "",
            "address_match_status": "unmatched", "address_match_confirmed_by": None,
            "address_match_confirmed_at": None,
            "address_match_candidates": [{"entry_id": 12, "score": .7, "method": "rule", "reason": "地址相符"}],
            "landlord_name": "不应导出", "identity_number": "不应导出", "phone": "不应导出"}
ENTRY = {"id": 12, "name": "合成小区", "community_id": 8, "community_name": "合成社区",
         "detail_address": "合成路", "aliases": ["合成旧名"], "address_type": "community"}


def annotated(properties=None, entries=None, mutate=None):
    workbook = load_workbook(xlsx.build_annotation_workbook(properties or [PROPERTY], entries or [ENTRY], USER["id"]))
    sheet = workbook["房屋标注"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, columns["decision"], "match")
        sheet.cell(row, columns["annotated_small_community_id"], 12)
        sheet.cell(row, columns["annotation_reason"], "正式名称与地址相符")
    if mutate:
        mutate(workbook, sheet, columns)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_export_is_flat_and_has_catalog_candidates_and_only_required_property_data():
    workbook = load_workbook(xlsx.build_annotation_workbook([PROPERTY, {**PROPERTY, "id": 41}], [ENTRY], 7))
    assert workbook.sheetnames == ["房屋标注", "小区地址库", "候选小区", "标注说明"]
    sheet = workbook["房屋标注"]
    assert sheet["A2"].value == "42" and sheet["A3"].value == "41"
    assert sheet.freeze_panes == "A2"
    assert list(sheet.values)[0] == xlsx.HEADERS
    assert not {"landlord_name", "identity_number", "phone"}.intersection(xlsx.HEADERS)
    assert workbook["小区地址库"]["F2"].value == '["合成旧名"]'
    assert workbook["候选小区"]["C2"].value == 12
    assert len(sheet.data_validations.dataValidation) == 1


def test_roundtrip_handles_literal_formula_prefixes_and_blank_catalog_fields():
    row = {**PROPERTY, "natural_address": "=合成地址"}
    entry = {**ENTRY, "detail_address": ""}
    content = annotated([row], [entry])
    rows, tokens = xlsx.parse_annotation_workbook(content, 7)
    assert not rows[0]["formula"]
    assert rows[0]["natural_address"] == "=合成地址"
    assert xlsx.valid_token(rows[0]["snapshot_token"], xlsx.property_token(row, 7))
    assert tokens[12] == xlsx.entry_token(entry, 7)


def test_catalog_tamper_and_wrong_export_account_are_rejected():
    with pytest.raises(ValueError, match="原始资料"):
        xlsx.parse_annotation_workbook(annotated(mutate=lambda workbook, *_: setattr(workbook["小区地址库"]["B2"], "value", "替换名称")), 7)
    with pytest.raises(ValueError, match="原导出账号"):
        xlsx.parse_annotation_workbook(annotated(), 99)


def test_invalid_workbook_and_non_ascii_signature_are_rejected():
    with pytest.raises(ValueError, match="无法读取"):
        xlsx.parse_annotation_workbook(b"not-an-xlsx", 7)
    assert not xlsx.valid_token("无效签名", "a" * 64)


def test_upload_limit_and_duplicate_records_and_formula_detection(monkeypatch):
    monkeypatch.setattr(xlsx, "MAX_UPLOAD_BYTES", 4)
    with pytest.raises(ValueError, match="10MB"):
        xlsx.parse_annotation_workbook(b"12345", 7)
    monkeypatch.setattr(xlsx, "MAX_UPLOAD_BYTES", 10 * 1024 * 1024)
    rows, _ = xlsx.parse_annotation_workbook(annotated([PROPERTY, PROPERTY]), 7)
    assert all(row["duplicate"] for row in rows)
    rows, _ = xlsx.parse_annotation_workbook(annotated(mutate=lambda _, sheet, columns: sheet.cell(2, columns["annotation_reason"], "=1+1")), 7)
    assert rows[0]["formula"]


class Connection:
    def __init__(self):
        self.sql = []
        self.params = []
        self.began = self.committed = self.rolled_back = False

    def cursor(self):
        return self

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_):
        return False

    async def execute(self, sql, params=()):
        self.sql.append(sql)
        self.params.append(params)

    async def executemany(self, sql, params):
        await self.execute(sql, params)

    async def fetchall(self):
        return [(42, 8, "合成社区", 3, "active")]

    async def begin(self):
        self.began = True

    async def commit(self):
        self.committed = True

    async def rollback(self):
        self.rolled_back = True


@pytest.fixture
def data(monkeypatch):
    properties = {42: dict(PROPERTY)}
    entries = [dict(ENTRY)]
    monkeypatch.setattr(registry, "_allowed_community_ids", AsyncMock(return_value=[8]))
    monkeypatch.setattr(registry, "_annotation_properties", AsyncMock(side_effect=lambda *_args, **_kw: properties))
    monkeypatch.setattr(registry, "_annotation_entries", AsyncMock(side_effect=lambda *_args, **_kw: entries))
    monkeypatch.setattr(registry, "record_admin_audit", AsyncMock())
    monkeypatch.setattr(registry, "request_audit_fields", lambda *_: {})
    return properties, entries


async def preview(content):
    return await registry.preview_property_annotations(None, UploadFile(file=BytesIO(content), filename="labels.xlsx"), USER, Connection())


@pytest.mark.asyncio
async def test_preview_is_read_only_and_apply_reuses_confirm_in_transaction(data, monkeypatch):
    result = await preview(annotated())
    assert result["ready"] == 1 and result["blocked"] == 0
    conn = Connection()
    monkeypatch.setattr(registry.db_manager, "get_pool", lambda *_: pytest.fail("must not acquire OnlineData connection"))
    applied = await registry.apply_property_annotations(
        registry.PropertyAnnotationApply(confirm=True, items=[result["items"][0]["apply_item"]]), None, USER, conn,
    )
    assert applied["confirmed"] == 1 and conn.committed
    assert any("current_version=current_version+1" in sql for sql in conn.sql)
    assert any("JSON_OBJECT('source','manual')" in sql for sql in conn.sql)
    assert all("natural_address=" not in sql for sql in conn.sql)
    assert registry._annotation_properties.call_args.kwargs["lock"]
    assert registry._annotation_entries.call_args.kwargs["lock"]


@pytest.mark.asyncio
@pytest.mark.parametrize("changed", [
    {"version": 4}, {"natural_address": "已变化"}, {"status": "inactive"}, {"household_status": "已注销"},
    {"address_match_status": "confirmed", "small_community_id": 15}, {"community_id": 9},
])
async def test_changed_snapshot_or_permission_blocks_preview(data, changed):
    data[0][42].update(changed)
    result = await preview(annotated())
    assert result["blocked"] == 1 and not result["items"][0].get("apply_item")


@pytest.mark.asyncio
async def test_legacy_inactive_cancelled_property_can_confirm_community(data):
    property_row = {**PROPERTY, "status": "inactive", "household_status": "已注销"}
    data[0][42] = property_row
    result = await preview(annotated([property_row]))
    assert result["ready"] == 1
    conn = Connection()
    conn.fetchall = AsyncMock(return_value=[(42, 8, "合成社区", 3, "inactive")])
    applied = await registry.apply_property_annotations(
        registry.PropertyAnnotationApply(confirm=True, items=[result["items"][0]["apply_item"]]), None, USER, conn,
    )
    assert applied["confirmed"] == 1 and conn.committed


@pytest.mark.asyncio
@pytest.mark.parametrize("changed", [{"community_id": 9}, {"name": "新名称"}, {"aliases": ["新别名"]}])
async def test_cross_community_and_changed_catalog_block_preview(data, changed):
    data[1][0].update(changed)
    assert (await preview(annotated()))["blocked"] == 1


@pytest.mark.asyncio
async def test_disabled_entry_missing_reason_and_duplicate_block_preview(data):
    data[1].clear()
    assert (await preview(annotated()))["blocked"] == 1
    data[1].append(dict(ENTRY))
    result = await preview(annotated(mutate=lambda _, sheet, columns: setattr(sheet.cell(2, columns["annotation_reason"]), "value", None)))
    assert result["blocked"] == 1
    assert (await preview(annotated([PROPERTY, PROPERTY])))["blocked"] == 2


@pytest.mark.asyncio
async def test_existing_manual_confirmation_is_marked_and_unchanged_result_is_skipped(data):
    manual = {**PROPERTY, "address_match_status": "confirmed", "small_community_id": 15}
    data[0][42] = manual
    result = await preview(annotated([manual]))
    assert result["items"][0]["replaces_manual"] is True
    manual["small_community_id"] = 12
    result = await preview(annotated([manual]))
    assert result["skipped"] == 1


@pytest.mark.asyncio
async def test_review_reason_is_visible_without_creating_apply_item(data):
    def review_row(_, sheet, columns):
        sheet.cell(2, columns["decision"], "review")
        sheet.cell(2, columns["annotated_small_community_id"]).value = None
    result = await preview(annotated(mutate=review_row))
    assert result["review"] == 1
    assert result["items"][0]["annotation_reason"] == "正式名称与地址相符"
    assert not result["items"][0].get("apply_item")


@pytest.mark.asyncio
async def test_apply_rejects_modified_preview_and_later_concurrent_change(data):
    item = (await preview(annotated()))["items"][0]["apply_item"]
    with pytest.raises(HTTPException) as error:
        await registry.apply_property_annotations(registry.PropertyAnnotationApply(confirm=True, items=[{**item, "small_community_id": 13}]), None, USER, Connection())
    assert error.value.status_code == 422
    data[0][42]["version"] = 4
    conn = Connection()
    with pytest.raises(HTTPException) as error:
        await registry.apply_property_annotations(registry.PropertyAnnotationApply(confirm=True, items=[item]), None, USER, conn)
    assert error.value.status_code == 409 and conn.rolled_back and not conn.committed
    assert not any("INSERT INTO" in sql for sql in conn.sql)


@pytest.mark.asyncio
@pytest.mark.parametrize("change", ["entry", "permission"])
async def test_apply_rechecks_catalog_and_permission_with_atomic_rollback(data, monkeypatch, change):
    item = (await preview(annotated()))["items"][0]["apply_item"]
    if change == "entry":
        data[1][0]["aliases"] = ["已变化别名"]
    else:
        monkeypatch.setattr(registry, "_allowed_community_ids", AsyncMock(return_value=[]))
    conn = Connection()
    with pytest.raises(HTTPException) as error:
        await registry.apply_property_annotations(registry.PropertyAnnotationApply(confirm=True, items=[item]), None, USER, conn)
    assert error.value.status_code in (403, 409)
    assert conn.rolled_back and not conn.committed
    assert not any("INSERT INTO" in sql for sql in conn.sql)


@pytest.mark.asyncio
async def test_export_reuses_search_and_sort_but_excludes_manage_scope(data, monkeypatch):
    monkeypatch.setattr(registry, "_property_search_result", AsyncMock(return_value={"data": [PROPERTY, {**PROPERTY, "id": 43, "community_id": 9}]}))
    response = await registry.export_property_annotations(registry.PropertySearch(sort="address_asc"), None, USER, Connection())
    assert response.media_type.endswith("sheet")
    assert registry._property_search_result.call_args.kwargs["export_all"] is True
    assert registry._property_search_result.call_args.args[0].sort == "address_asc"
    audit = registry.record_admin_audit.call_args.kwargs["detail"]
    assert audit["rows"] == 1 and "合成路" not in str(audit)


def test_routes_require_property_manage_and_precede_dynamic_route():
    paths = [route.path for route in registry.router.routes]
    for action in ("export", "preview", "apply"):
        path = "/api/registry/properties/small-community-annotations/" + action
        assert paths.index(path) < paths.index("/api/registry/properties/{property_id}/people")
        route = next(route for route in registry.router.routes if route.path == path)
        dependency = next(dep.call for dep in route.dependant.dependencies if dep.name == "user")
        assert "registry.property.manage" in str([cell.cell_contents for cell in dependency.__closure__])


@pytest.mark.asyncio
@pytest.mark.parametrize("addresses_active,platform_active", [(False, False), (True, True), (True, False), (False, True)])
async def test_catalog_uses_current_domain_tables_scope_and_locks(monkeypatch, addresses_active, platform_active):
    monkeypatch.setattr(registry.settings, "REGISTRY_ADDRESS_DOMAIN_ACTIVE", addresses_active)
    monkeypatch.setattr(registry.settings, "PLATFORM_DOMAIN_ACTIVE", platform_active)
    cur = Connection()
    cur.fetchall = AsyncMock(return_value=[(12, "合成小区", 8, "合成社区", "合成路", '["合成旧名"]', "community")])
    entries = await registry._annotation_entries(cur, community_ids=[8], entry_ids=[12], lock=True)
    address_schema = registry.settings.MYSQL_REGISTRY_DB if addresses_active else registry.settings.MYSQL_ONLINE_DATA_DB
    community_schema = registry.settings.MYSQL_PLATFORM_DB if platform_active else registry.settings.MYSQL_ONLINE_DATA_DB
    assert f"FROM `{address_schema}`._police_address_entries" in cur.sql[0]
    assert f"JOIN `{community_schema}`._communities" in cur.sql[0]
    assert "entry.enabled=1 AND community.is_active=1" in cur.sql[0]
    assert cur.sql[0].endswith("FOR UPDATE")
    assert cur.params == [(12, 8)]
    assert entries == [ENTRY]


@pytest.mark.asyncio
async def test_property_snapshot_loading_matches_export_serialization():
    from datetime import datetime
    cur = Connection()
    values = [PROPERTY[key] for key in xlsx.SNAPSHOT_FIELDS]
    values[10] = datetime.fromisoformat(values[10])
    cur.fetchall = AsyncMock(return_value=[tuple(values)])
    rows = await registry._annotation_properties(cur, [42], lock=True)
    assert xlsx.property_token(rows[42], 7) == xlsx.property_token(PROPERTY, 7)
    assert cur.sql[0].endswith("FOR UPDATE")
